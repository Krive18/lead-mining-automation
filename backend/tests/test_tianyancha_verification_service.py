import base64
import io
import tempfile
import unittest
from pathlib import Path
from threading import Event, Thread
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook

from app.model.schema import CollectedLeadDTO
from app.service.tianyancha_verification_service import (
    TIANYANCHA_EXPORT_HEADER_LIST,
    TIANYANCHA_STATUS_MATCHED,
    TIANYANCHA_STATUS_NOT_FOUND,
    TIANYANCHA_STATUS_UNCONFIRMED,
    TianyanchaVerificationService,
)


class FakeTianyanchaCollector:
    def __init__(self, lead_map=None, error_on_call=None, block_on_call=None):
        self.lead_map = lead_map or {}
        self.error_on_call = error_on_call
        self.block_on_call = block_on_call
        self.call_list = []
        self.callback = None
        self.stop_requested_callback = lambda: False
        self.blocked_event = Event()

    def set_task_event_callback(self, callback):
        self.callback = callback

    def set_stop_requested_callback(self, callback):
        self.stop_requested_callback = callback

    def collect_by_company_names(self, industry, company_name_list, collect_limit):
        self.call_list.append(list(company_name_list))
        if self.error_on_call is not None and len(self.call_list) == self.error_on_call:
            raise RuntimeError("天眼查页面出现登录/验证/风控要求")
        if self.block_on_call is not None and len(self.call_list) == self.block_on_call:
            self.blocked_event.set()
            while not self.stop_requested_callback():
                self.blocked_event.wait(0.01)
            raise RuntimeError("用户请求停止天眼查反证")
        result_list = []
        for company_name in company_name_list:
            lead = self.lead_map.get(company_name)
            if lead is not None:
                result_list.append(lead)
        return result_list[:collect_limit]


def build_excel_base64(header_list, row_list):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(header_list)
    for row in row_list:
        sheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    return base64.b64encode(output.getvalue()).decode("ascii")


def build_lead(company_name, phone="", unified_social_credit_code=""):
    return CollectedLeadDTO(
        source_type="tianyancha_public_web",
        source_keyword=company_name,
        industry="",
        suspected_company_name=company_name,
        legal_person="张三",
        contact_name="张三",
        phone=phone,
        phone_source="天眼查公开页可见字段" if phone else "",
        company_status="存续",
        registered_address="广东省广州市天河区测试路1号",
        business_scope="服装销售；互联网销售",
        unified_social_credit_code=unified_social_credit_code,
        profile_url=f"https://www.tianyancha.com/company/{company_name}",
        is_enterprise_verified=True,
    )


class TianyanchaVerificationServiceTest(unittest.TestCase):
    def test_export_uses_compact_standard_columns_and_separates_max_shareholder(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            service = TianyanchaVerificationService(
                collector_factory=lambda: FakeTianyanchaCollector({}),
                config=SimpleNamespace(export_dir=Path(temp_dir)),
                sleep_func=lambda seconds: None,
                random_uniform_func=lambda left, right: 0.0,
            )
            lead = build_lead("广州测试科技有限公司", phone="13800138000")
            lead.legal_person = "张三"
            lead.contact_name = "法人：张三；最大股东：李四"
            result = service.build_result_from_lead("广州测试科技有限公司", lead)

            workbook_bytes = service.write_result_workbook(
                ["公司", "原表备注"],
                [
                    {
                        "company_name": "广州测试科技有限公司",
                        "row_values": ["广州测试科技有限公司", "不应导出"],
                    }
                ],
                [result],
            )
            workbook = load_workbook(io.BytesIO(workbook_bytes), read_only=True, data_only=True)
            try:
                rows = list(workbook["全部结果"].iter_rows(values_only=True))
                self.assertEqual(list(rows[0]), TIANYANCHA_EXPORT_HEADER_LIST)
                self.assertEqual(rows[1][0], "广州测试科技有限公司")
                self.assertEqual(rows[1][4], "张三")
                self.assertEqual(rows[1][5], "李四")
                self.assertNotIn("原表备注", rows[0])
                self.assertNotIn("不应导出", rows[1])
            finally:
                workbook.close()

    def test_preview_and_run_job_with_batch_jitter_and_export(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            fake_collector = FakeTianyanchaCollector(
                {
                    "广州测试服装有限公司": build_lead(
                        "广州测试服装有限公司",
                        "13800138000",
                        "91440101MA5CKX7R2L",
                    ),
                    "佛山测试科技有限公司": build_lead("佛山测试科技有限公司"),
                }
            )
            sleep_list = []
            service = TianyanchaVerificationService(
                collector_factory=lambda: fake_collector,
                config=SimpleNamespace(export_dir=Path(temp_dir)),
                sleep_func=sleep_list.append,
                random_uniform_func=lambda left, right: 2.0,
            )
            excel_base64 = build_excel_base64(
                ["公司", "备注"],
                [
                    ["广州测试服装有限公司", "A"],
                    ["深圳不存在有限公司", "B"],
                    ["佛山测试科技有限公司", "C"],
                ],
            )

            preview = service.preview_excel_columns_base64(excel_base64, "company.xlsx")
            self.assertEqual(preview["suggested_column_index"], 1)

            snapshot = service.create_verification_job(
                excel_base64,
                "company.xlsx",
                company_name_column_index=1,
                batch_size=2,
                jitter_min_seconds=2,
                jitter_max_seconds=2,
                batch_cooldown_seconds=3,
            )
            service.run_verification_job(snapshot["job_id"])
            result = service.get_verification_job_snapshot(snapshot["job_id"])

            self.assertEqual(result["status"], "completed")
            self.assertEqual(result["matched_count"], 2)
            self.assertEqual(result["not_found_count"], 1)
            self.assertEqual(result["unconfirmed_count"], 0)
            self.assertEqual(fake_collector.call_list, [["广州测试服装有限公司", "深圳不存在有限公司"], ["佛山测试科技有限公司"]])
            self.assertEqual(sleep_list, [5.0])
            self.assertTrue(Path(result["file_path"]).exists())

            workbook = load_workbook(result["file_path"], read_only=True, data_only=True)
            try:
                rows = list(workbook["全部结果"].iter_rows(values_only=True))
                self.assertEqual(list(rows[0]), TIANYANCHA_EXPORT_HEADER_LIST)
                status_column = rows[0].index("天眼查反证状态")
                phone_column = rows[0].index("联系电话")
                self.assertEqual(rows[1][status_column], TIANYANCHA_STATUS_MATCHED)
                self.assertEqual(rows[2][status_column], TIANYANCHA_STATUS_NOT_FOUND)
                self.assertEqual(rows[1][phone_column], "13800138000")
            finally:
                workbook.close()

    def test_batch_error_keeps_rows_as_unconfirmed(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            fake_collector = FakeTianyanchaCollector(
                {"广州测试服装有限公司": build_lead("广州测试服装有限公司")},
                error_on_call=2,
            )
            service = TianyanchaVerificationService(
                collector_factory=lambda: fake_collector,
                config=SimpleNamespace(export_dir=Path(temp_dir)),
                sleep_func=lambda seconds: None,
                random_uniform_func=lambda left, right: 0.0,
            )
            excel_base64 = build_excel_base64(
                ["公司"],
                [
                    ["广州测试服装有限公司"],
                    ["佛山风控有限公司"],
                ],
            )

            snapshot = service.create_verification_job(
                excel_base64,
                "company.xlsx",
                company_name_column_index=1,
                batch_size=1,
                jitter_min_seconds=0,
                jitter_max_seconds=0,
                batch_cooldown_seconds=0,
            )
            service.run_verification_job(snapshot["job_id"])
            result = service.get_verification_job_snapshot(snapshot["job_id"])

            self.assertEqual(result["status"], "completed")
            self.assertEqual(result["matched_count"], 1)
            self.assertEqual(result["unconfirmed_count"], 1)
            self.assertEqual(result["result_items"][1]["status"], TIANYANCHA_STATUS_UNCONFIRMED)
            self.assertIn("风控", result["result_items"][1]["remark"])


    def test_latest_snapshot_recovers_completed_download_and_events(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            fake_collector = FakeTianyanchaCollector(
                {"骞垮窞娴嬭瘯鏈嶈鏈夐檺鍏徃": build_lead("骞垮窞娴嬭瘯鏈嶈鏈夐檺鍏徃")}
            )
            service = TianyanchaVerificationService(
                collector_factory=lambda: fake_collector,
                config=SimpleNamespace(export_dir=Path(temp_dir)),
                sleep_func=lambda seconds: None,
                random_uniform_func=lambda left, right: 0.0,
            )
            self.assertEqual(service.get_latest_verification_job_snapshot()["status"], "none")
            excel_base64 = build_excel_base64(["鍏徃"], [["骞垮窞娴嬭瘯鏈嶈鏈夐檺鍏徃"]])

            snapshot = service.create_verification_job(
                excel_base64,
                "company.xlsx",
                company_name_column_index=1,
            )
            self.assertEqual(service.get_latest_verification_job_snapshot()["job_id"], snapshot["job_id"])
            service.run_verification_job(snapshot["job_id"])
            latest = service.get_latest_verification_job_snapshot()

            self.assertEqual(latest["status"], "completed")
            self.assertTrue(latest["download_url"])
            self.assertTrue(latest["event_items"])

    def test_history_lists_recent_jobs_without_result_rows(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            fake_collector = FakeTianyanchaCollector({})
            service = TianyanchaVerificationService(
                collector_factory=lambda: fake_collector,
                config=SimpleNamespace(export_dir=Path(temp_dir)),
                sleep_func=lambda seconds: None,
                random_uniform_func=lambda left, right: 0.0,
            )
            first_excel = build_excel_base64(["公司"], [["广州第一服装有限公司"]])
            second_excel = build_excel_base64(["公司"], [["广州第二服装有限公司"]])

            first = service.create_verification_job(first_excel, "first.xlsx", company_name_column_index=1)
            second = service.create_verification_job(second_excel, "second.xlsx", company_name_column_index=1)
            service.run_verification_job(first["job_id"])
            service.run_verification_job(second["job_id"])

            history = service.list_verification_job_history(limit=10)

            self.assertEqual([item["job_id"] for item in history], [second["job_id"], first["job_id"]])
            self.assertEqual(history[0]["file_name"], "second.xlsx")
            self.assertEqual(history[0]["total_count"], 1)
            self.assertIn("created_at", history[0])
            self.assertIn("updated_at", history[0])
            self.assertNotIn("result_items", history[0])
            self.assertNotIn("event_items", history[0])

    def test_stop_job_preserves_completed_batches_and_exports_partial_result(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            fake_collector = FakeTianyanchaCollector(
                {"广州已完成有限公司": build_lead("广州已完成有限公司")},
                block_on_call=2,
            )
            service = TianyanchaVerificationService(
                collector_factory=lambda: fake_collector,
                config=SimpleNamespace(export_dir=Path(temp_dir)),
                sleep_func=lambda seconds: None,
                random_uniform_func=lambda left, right: 0.0,
            )
            excel_base64 = build_excel_base64(
                ["公司"],
                [["广州已完成有限公司"], ["广州等待停止有限公司"]],
            )
            snapshot = service.create_verification_job(
                excel_base64,
                "stop-company.xlsx",
                company_name_column_index=1,
                batch_size=1,
                jitter_min_seconds=0,
                jitter_max_seconds=0,
                batch_cooldown_seconds=0,
            )

            run_thread = Thread(target=service.run_verification_job, args=(snapshot["job_id"],), daemon=True)
            run_thread.start()
            self.assertTrue(fake_collector.blocked_event.wait(2))
            service.stop_verification_job(snapshot["job_id"])
            run_thread.join(2)
            result = service.get_verification_job_snapshot(snapshot["job_id"])

            self.assertFalse(run_thread.is_alive())
            self.assertEqual(result["status"], "stopped")
            self.assertEqual(result["processed_count"], 1)
            self.assertEqual(result["matched_count"], 1)
            self.assertTrue(result["download_url"])
            self.assertTrue(Path(result["file_path"]).exists())


if __name__ == "__main__":
    unittest.main()
