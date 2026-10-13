import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

from openpyxl import load_workbook
from pydantic import ValidationError

from app.core.constant import SOURCE_TYPE_ALIBABA_1688
from app.model.schema import ExportCreateDTO
from app.repository.lead_repository import LeadRepository
from app.service.export_service import EXPORT_HEADER, ExportService


class ExportServiceTest(unittest.TestCase):
    def test_alibaba_1688_export_source_homepage_never_uses_tianyancha_url(self):
        service = ExportService(repository=None)
        lead = {
            "id": "lead-1",
            "source_type": SOURCE_TYPE_ALIBABA_1688,
            "source_keyword": "\u6df1\u5733\u624b\u673a\u6570\u636e\u7ebf\u5de5\u5382",
            "shop_url": "",
            "profile_url": "https://www.tianyancha.com/company/123",
            "assignment_status": "assigned",
        }

        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "export.xlsx"
            service.write_workbook(file_path, [lead], is_phone_masked=False)
            workbook = load_workbook(file_path)
            sheet = workbook.active
            profile_column = [name for name, _ in EXPORT_HEADER].index("profile_url") + 1
            source_homepage = sheet.cell(row=2, column=profile_column).value

        self.assertTrue(str(source_homepage).startswith("https://s.1688.com/"))
        self.assertNotIn("tianyancha", str(source_homepage))

    def test_export_cleans_registered_address_before_writing_excel(self):
        service = ExportService(repository=None)
        lead = {
            "id": "lead-1",
            "source_type": SOURCE_TYPE_ALIBABA_1688,
            "registered_address": (
                "\u5e7f\u5dde\u5e02\u767d\u4e91\u533a\u9e64\u9f99\u8857"
                "\u9e64\u8fb9\u901a\u8fbe\u521b\u610f\u56ed\u9644\u8fd1\u516c\u53f8"
                "\u7ecf\u8425\u8303\u56f4\u670d\u88c5\u5236\u9020;\u5316\u5986\u54c1\u6279\u53d1"
            ),
            "assignment_status": "assigned",
        }

        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "export.xlsx"
            service.write_workbook(file_path, [lead], is_phone_masked=False)
            workbook = load_workbook(file_path)
            sheet = workbook.active
            address_column = [name for name, _ in EXPORT_HEADER].index("registered_address") + 1
            address = sheet.cell(row=2, column=address_column).value

        self.assertEqual(address, "\u5e7f\u5dde\u5e02\u767d\u4e91\u533a\u9e64\u9f99\u8857\u9e64\u8fb9\u901a\u8fbe\u521b\u610f\u56ed")

    def test_create_export_passes_industry_filter_to_repository(self):
        repository = MagicMock()
        repository.list_lead_for_export.return_value = []
        service = ExportService(repository=repository)

        with tempfile.TemporaryDirectory() as temp_dir, patch(
            "app.service.export_service.get_config",
            return_value=SimpleNamespace(export_dir=Path(temp_dir)),
        ):
            service.create_export(
                ExportCreateDTO(
                    levels=["A", "B"],
                    industry="\u7f8e\u5986\u4e2a\u62a4",
                )
            )

        repository.list_lead_for_export.assert_called_once()
        self.assertEqual(repository.list_lead_for_export.call_args.kwargs["industry"], "\u7f8e\u5986\u4e2a\u62a4")

    def test_create_export_filters_and_names_file_by_source_group(self):
        repository = MagicMock()
        repository.list_lead_for_export.return_value = []
        service = ExportService(repository=repository)

        with tempfile.TemporaryDirectory() as temp_dir, patch(
            "app.service.export_service.get_config",
            return_value=SimpleNamespace(export_dir=Path(temp_dir)),
        ):
            result = service.create_export(
                ExportCreateDTO(
                    levels=["A", "B", "C", "D"],
                    source_group="alibaba_international",
                )
            )

        repository.list_lead_for_export.assert_called_once()
        self.assertEqual(
            repository.list_lead_for_export.call_args.kwargs["source_group"],
            "alibaba_international",
        )
        self.assertTrue(result["file_name"].startswith("\u963f\u91cc\u56fd\u9645\u7ad9_\u7ebf\u7d22\u5bfc\u51fa_"))

    def test_export_create_dto_rejects_unknown_source_group(self):
        with self.assertRaises(ValidationError):
            ExportCreateDTO(source_group="unknown-source-group")

    def test_repository_export_where_clause_filters_by_industry(self):
        repository = LeadRepository()

        where_sql, parameters = repository.build_lead_where_clause(industry="3C\u6570\u7801")

        self.assertIn("industry = ?", where_sql)
        self.assertIn("3C\u6570\u7801", parameters)

    def test_export_writes_unified_social_credit_code_column(self):
        service = ExportService(repository=None)
        lead = {
            "id": "lead-credit-code",
            "source_type": SOURCE_TYPE_ALIBABA_1688,
            "unified_social_credit_code": "91440101MA5CKX7R2L",
            "assignment_status": "assigned",
        }

        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "credit-code.xlsx"
            service.write_workbook(file_path, [lead], is_phone_masked=False)
            workbook = load_workbook(file_path, read_only=True, data_only=True)
            try:
                sheet = workbook.active
                code_column = [name for name, _ in EXPORT_HEADER].index("unified_social_credit_code") + 1
                self.assertEqual(sheet.cell(row=1, column=code_column).value, "统一社会信用代码")
                self.assertEqual(sheet.cell(row=2, column=code_column).value, "91440101MA5CKX7R2L")
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
