import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import load_workbook
from pydantic import ValidationError

from app.core.database import init_database
from app.model.schema import CollectedLeadDTO, TaskCreateDTO
from app.repository.lead_repository import LeadRepository
from app.service.export_service import EXPORT_HEADER, ExportService


SOURCE_TYPE_ALIBABA_INTERNATIONAL = "alibaba_international"
COMPANY_NAME_EN_HEADER = "英文公司名称"
SHOP_LEVEL_HEADER = "店铺等级"
MYSQL_COMPANY_NAME_EN_COLUMN = "`company_name_en` VARCHAR(255) NOT NULL DEFAULT ''"
MYSQL_SHOP_LEVEL_COLUMN = "`shop_level` VARCHAR(255) NOT NULL DEFAULT ''"


class AlibabaInternationalFieldChainTest(unittest.TestCase):
    def test_task_create_accepts_alibaba_international_source_type(self):
        try:
            request = TaskCreateDTO(
                source_type=SOURCE_TYPE_ALIBABA_INTERNATIONAL,
                industry="3C数码",
                keywords=["data cable factory"],
                collect_limit=5,
                is_qcc_enabled=False,
                qcc_limit=0,
            )
        except ValidationError as exc:
            self.fail(f"TaskCreateDTO should accept alibaba_international: {exc}")

        self.assertEqual(request.source_type, SOURCE_TYPE_ALIBABA_INTERNATIONAL)

    def test_collected_lead_accepts_alibaba_international_source_type(self):
        try:
            lead = CollectedLeadDTO(
                source_type=SOURCE_TYPE_ALIBABA_INTERNATIONAL,
                source_keyword="data cable factory",
                industry="3C数码",
                suspected_company_name="深圳市示例科技有限公司",
                company_name_en="Shenzhen Example Technology Co., Ltd.",
                shop_level="Gold Supplier",
            )
        except ValidationError as exc:
            self.fail(f"CollectedLeadDTO should accept alibaba_international: {exc}")

        self.assertEqual(lead.source_type, SOURCE_TYPE_ALIBABA_INTERNATIONAL)
        self.assertEqual(lead.company_name_en, "Shenzhen Example Technology Co., Ltd.")
        self.assertEqual(lead.shop_level, "Gold Supplier")

    def test_repository_persists_and_reads_company_name_en_and_shop_level(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            config = SimpleNamespace(
                database_driver="sqlite",
                database_path=Path(temp_dir) / "alibaba-international.db",
            )
            with patch("app.core.database.get_config", return_value=config):
                init_database()
                repository = LeadRepository()
                repository.insert_lead(
                    {
                        "id": "lead-alibaba-international",
                        "task_id": "task-alibaba-international",
                        "source_type": SOURCE_TYPE_ALIBABA_INTERNATIONAL,
                        "source_keyword": "data cable factory",
                        "industry": "3C数码",
                        "suspected_company_name": "深圳市示例科技有限公司",
                        "company_name": "深圳市示例科技有限公司",
                        "company_name_en": "Shenzhen Example Technology Co., Ltd.",
                        "douyin_nickname": "",
                        "product_name": "USB data cable",
                        "shop_level": "Gold Supplier",
                        "contact_name": "",
                        "phone": "",
                        "phone_source": "",
                        "website_url": "",
                        "shop_url": "https://example.en.alibaba.com",
                        "profile_url": "",
                        "sample_video_url": "",
                        "is_blue_v": 0,
                        "is_enterprise_verified": 1,
                        "company_status": "存续",
                        "unified_social_credit_code": "",
                        "legal_person": "",
                        "business_scope": "USB data cable manufacturing",
                        "registered_address": "Guangdong Shenzhen",
                        "province": "广东省",
                        "city": "深圳市",
                        "district": "南山区",
                        "city_source": "manual",
                        "city_confidence": "high",
                        "risk_tag": "",
                        "verification_status": "public_web",
                        "confidence": "high",
                        "score": 80,
                        "level": "A",
                        "assignment_status": "unassigned",
                        "gmt_create": "2026-07-13 10:00:00",
                        "gmt_modified": "2026-07-13 10:00:00",
                    }
                )

                lead = repository.get_lead_by_id("lead-alibaba-international")

        self.assertEqual(lead.get("company_name_en"), "Shenzhen Example Technology Co., Ltd.")
        self.assertEqual(lead.get("shop_level"), "Gold Supplier")

    def test_export_uses_exact_unicode_header_literals(self):
        field_names = [field_name for field_name, _ in EXPORT_HEADER]
        self.assertIn("company_name_en", field_names)
        self.assertIn("shop_level", field_names)
        self.assertEqual(field_names[field_names.index("company_name") + 1], "company_name_en")
        self.assertEqual(field_names[field_names.index("product_name") + 1], "shop_level")
        self.assertIn(("company_name_en", COMPANY_NAME_EN_HEADER), EXPORT_HEADER)
        self.assertIn(("shop_level", SHOP_LEVEL_HEADER), EXPORT_HEADER)

    def test_export_writes_exact_unicode_headers_and_values(self):
        service = ExportService(repository=None)
        lead = {
            "id": "lead-export-alibaba-international",
            "source_type": SOURCE_TYPE_ALIBABA_INTERNATIONAL,
            "company_name": "深圳市示例科技有限公司",
            "company_name_en": "Shenzhen Example Technology Co., Ltd.",
            "product_name": "USB data cable",
            "shop_level": "Gold Supplier",
            "assignment_status": "assigned",
        }

        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "alibaba-international.xlsx"
            service.write_workbook(file_path, [lead], is_phone_masked=False)
            workbook = load_workbook(file_path, read_only=True, data_only=True)
            try:
                sheet = workbook.active
                headers = [cell.value for cell in sheet[1]]
                company_name_en_column = headers.index(COMPANY_NAME_EN_HEADER) + 1
                shop_level_column = headers.index(SHOP_LEVEL_HEADER) + 1

                self.assertEqual(sheet.cell(row=2, column=company_name_en_column).value, lead["company_name_en"])
                self.assertEqual(sheet.cell(row=2, column=shop_level_column).value, lead["shop_level"])
            finally:
                workbook.close()

    def test_mysql_init_sql_includes_company_name_en_and_shop_level_columns(self):
        init_sql_path = Path(__file__).resolve().parents[2] / "database" / "init.sql"
        init_sql_text = init_sql_path.read_text(encoding="utf-8")

        self.assertIn(MYSQL_COMPANY_NAME_EN_COLUMN, init_sql_text)
        self.assertIn(MYSQL_SHOP_LEVEL_COLUMN, init_sql_text)


if __name__ == "__main__":
    unittest.main()
