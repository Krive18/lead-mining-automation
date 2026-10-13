import base64
import io
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook

from app.model.schema import CollectedLeadDTO
from app.service.qixin_api_service import QixinApiError, QixinCompanyResult
from app.service.qixin_verification_service import QixinVerificationService


def build_excel_base64(header_list, row_list):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(header_list)
    for row in row_list:
        sheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    return base64.b64encode(output.getvalue()).decode("ascii")


def build_qixin_lead(company_name):
    return CollectedLeadDTO(
        source_type="qixin_public_web",
        source_keyword=company_name,
        industry="",
        suspected_company_name=company_name,
        legal_person="张三",
        contact_name="法人：张三；最大股东：李四",
        phone="13812349876",
        phone_source="启信宝1.51企业联系方式",
        company_status="存续",
        unified_social_credit_code="91440300MA5F12345X",
        registered_address="广东省深圳市南山区测试路1号",
        business_scope="电子产品销售",
        profile_url="https://www.qixin.com/search?key=test",
        is_enterprise_verified=True,
    )


class FakeQixinApiService:
    def __init__(self):
        self.call_list = []

    def assert_configured(self):
        return None

    def query_company(self, company_name, industry):
        self.call_list.append(company_name)
        if "未授权" in company_name:
            raise QixinApiError(
                "未添加IP白名单",
                endpoint="/enterprise/getBasicInfo",
                status="104",
            )
        return QixinCompanyResult(
            lead=build_qixin_lead(company_name),
            endpoint_status={
                "basic": {"success": True, "authorized": True},
                "contact": {"success": True, "authorized": True},
                "partners": {"success": True, "authorized": True},
            },
        )


class QixinVerificationServiceTest(unittest.TestCase):
    def test_no_result_response_is_classified_as_not_found(self):
        service = QixinVerificationService(
            qixin_api_service=FakeQixinApiService(),
            config=SimpleNamespace(export_dir=Path(tempfile.gettempdir())),
            sleep_func=lambda seconds: None,
        )

        result = service.build_qixin_error_result(
            "广东天天美妆有限公司",
            QixinApiError("查询无结果", endpoint="/enterprise/getBasicInfo", status="201"),
        )

        self.assertEqual(result.status, "未找到")

    def test_upload_company_excel_queries_api_and_exports_complete_fields(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            fake_api_service = FakeQixinApiService()
            service = QixinVerificationService(
                qixin_api_service=fake_api_service,
                config=SimpleNamespace(export_dir=Path(temp_dir)),
                sleep_func=lambda seconds: None,
            )
            excel_base64 = build_excel_base64(
                ["公司名称", "备注"],
                [
                    ["深圳市测试科技有限公司", "A"],
                    ["深圳市未授权科技有限公司", "B"],
                ],
            )

            preview = service.preview_excel_columns_base64(excel_base64, "company.xlsx")
            self.assertEqual(preview["suggested_column_index"], 1)
            snapshot = service.create_verification_job(
                excel_base64,
                "company.xlsx",
                company_name_column_index=1,
                request_interval_seconds=0,
            )
            service.run_verification_job(snapshot["job_id"])
            result = service.get_verification_job_snapshot(snapshot["job_id"])

            self.assertEqual(result["status"], "completed")
            self.assertEqual(result["matched_count"], 1)
            self.assertEqual(result["unconfirmed_count"], 1)
            self.assertEqual(
                fake_api_service.call_list,
                ["深圳市测试科技有限公司", "深圳市未授权科技有限公司"],
            )
            self.assertTrue(Path(result["file_path"]).exists())

            workbook = load_workbook(result["file_path"], read_only=True, data_only=True)
            try:
                rows = list(workbook["全部结果"].iter_rows(values_only=True))
                self.assertIn("启信宝反证状态", rows[0])
                self.assertIn("统一社会信用代码", rows[0])
                self.assertIn("联系电话", rows[0])
                self.assertIn("91440300MA5F12345X", rows[1])
                self.assertIn("13812349876", rows[1])
                self.assertEqual(rows[2][rows[0].index("启信宝反证状态")], "未确认")
            finally:
                workbook.close()

    def test_stop_before_run_exports_empty_partial_result(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            service = QixinVerificationService(
                qixin_api_service=FakeQixinApiService(),
                config=SimpleNamespace(export_dir=Path(temp_dir)),
                sleep_func=lambda seconds: None,
            )
            excel_base64 = build_excel_base64(["公司名称"], [["深圳市测试科技有限公司"]])
            snapshot = service.create_verification_job(
                excel_base64,
                "stop.xlsx",
                company_name_column_index=1,
                request_interval_seconds=0,
            )

            service.stop_verification_job(snapshot["job_id"])
            service.run_verification_job(snapshot["job_id"])
            result = service.get_verification_job_snapshot(snapshot["job_id"])

            self.assertEqual(result["status"], "stopped")
            self.assertEqual(result["processed_count"], 0)
            self.assertTrue(Path(result["file_path"]).exists())


if __name__ == "__main__":
    unittest.main()
