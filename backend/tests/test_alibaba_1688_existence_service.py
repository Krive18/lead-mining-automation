import io
import base64
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook
from openpyxl import load_workbook

from app.service.alibaba_1688_existence_service import (
    ALIBABA_1688_STATUS_EXISTS,
    ALIBABA_1688_STATUS_NOT_EXISTS,
    ALIBABA_1688_STATUS_UNCONFIRMED,
    Alibaba1688ExistenceService,
    BingSearchIndexProvider,
    BrowserSearchIndexProvider,
    SearchIndexResult,
    clean_1688_company_name_preserving_full_name,
)


class FailingCollector:
    def build_search_url_list(self, keyword, page_number):
        return [f"https://s.1688.com/company/pc/factory_search.htm?keywords={keyword}&beginPage={page_number}"]

    def goto_public_page(self, page, url):
        raise RuntimeError("1688公开页面出现验证要求，已停止采集")


class DeterministicCollector:
    def __init__(self):
        self.current_url = ""

    def build_search_url_list(self, keyword, page_number):
        return [f"https://s.1688.com/company/pc/factory_search.htm?keywords={keyword}&beginPage={page_number}"]

    def goto_public_page(self, page, url):
        self.current_url = url

    def scroll_result_page_to_load(self, page):
        return None

    def extract_raw_item_list(self, page):
        if "存在公司" in self.current_url:
            return [{"company_name": "广州市存在公司有限公司"}]
        return [{"company_name": "广州市其他企业有限公司"}]


class FakeSearchIndexProvider:
    source_name = "Bing/微软搜索"

    def __init__(self, result_map=None, error=None):
        self.result_map = result_map or {}
        self.error = error
        self.call_list = []

    def search(self, query, limit=10):
        self.call_list.append((query, limit))
        if self.error is not None:
            raise self.error
        return self.result_map.get(query, [])

    def build_search_url(self, query):
        return f"https://cn.bing.com/search?q={query}"


class Alibaba1688ExistenceServiceTest(unittest.TestCase):
    def test_browser_provider_builds_results_from_search_page_anchors(self):
        result_list = BrowserSearchIndexProvider.build_result_list_from_anchor_list(
            [
                {
                    "title": "广东袋鼠妈妈生物科技有限公司 阿里巴巴店铺",
                    "url": "https://shop123456.1688.com/",
                    "snippet": "搜索结果页里出现的1688链接。",
                },
                {
                    "title": "广东袋鼠妈妈生物科技有限公司 - 企查查",
                    "url": "https://www.qcc.com/firm/example.html",
                    "snippet": "非1688结果。",
                },
            ],
            limit=10,
        )

        self.assertEqual(len(result_list), 2)
        self.assertEqual(result_list[0].url, "https://shop123456.1688.com/")
        self.assertIn("阿里巴巴店铺", result_list[0].title)

    def test_clean_company_name_preserves_prefix_before_parenthesized_region(self):
        company_name = "袋鼠（广东）生物科技有限公司"

        clean_name = clean_1688_company_name_preserving_full_name(company_name)

        self.assertEqual(clean_name, company_name)

    def test_bing_provider_parses_rss_results(self):
        rss_text = """<?xml version="1.0" encoding="utf-8" ?>
<rss version="2.0">
  <channel>
    <item>
      <title>广州庄雅生物科技有限公司-企业信息查询黄页-阿里巴巴</title>
      <link>https://www.1688.com/factory/example.html</link>
      <description>阿里巴巴提供广州庄雅生物科技有限公司主营产品信息。</description>
    </item>
  </channel>
</rss>"""

        result_list = BingSearchIndexProvider().parse_rss_result_list(rss_text)

        self.assertEqual(len(result_list), 1)
        self.assertEqual(result_list[0].title, "广州庄雅生物科技有限公司-企业信息查询黄页-阿里巴巴")
        self.assertEqual(result_list[0].url, "https://www.1688.com/factory/example.html")
        self.assertIn("主营产品", result_list[0].snippet)

    def test_judge_company_result_requires_exact_full_company_name(self):
        service = Alibaba1688ExistenceService(collector=object())

        result = service.judge_company_result(
            "广州市白云区点绘化妆品有限公司",
            [
                "广州点绘化妆品厂",
                "广州市白云区点绘化妆品有限公司",
            ],
            "https://example.test/search",
            "广州市白云区点绘化妆品有限公司",
        )

        self.assertEqual(result.status, ALIBABA_1688_STATUS_EXISTS)
        self.assertEqual(result.matched_company_name, "广州市白云区点绘化妆品有限公司")
        self.assertEqual(result.search_url, "https://example.test/search")
        self.assertEqual(result.search_keyword, "广州市白云区点绘化妆品有限公司")

    def test_judge_company_result_rejects_similar_short_name(self):
        service = Alibaba1688ExistenceService(collector=object())

        result = service.judge_company_result(
            "广州市白云区点绘化妆品有限公司",
            ["广州点绘化妆品厂", "广州市点绘化妆品商行"],
            "https://example.test/search",
            "广州市白云区点绘化妆品有限公司",
        )

        self.assertEqual(result.status, ALIBABA_1688_STATUS_NOT_EXISTS)
        self.assertEqual(result.matched_company_name, "")
        self.assertIn("严格匹配", result.remark)
        self.assertEqual(result.search_keyword, "广州市白云区点绘化妆品有限公司")

    def test_verify_company_row_records_access_limit_as_unconfirmed(self):
        service = Alibaba1688ExistenceService(
            collector=FailingCollector(),
            search_index_provider=FakeSearchIndexProvider(error=RuntimeError("搜索源出现验证或访问限制")),
        )

        result = service.verify_company_row_with_fallback(
            {"company_name": "广州市白云区点绘化妆品有限公司"},
            page=None,
        )

        self.assertEqual(result.status, ALIBABA_1688_STATUS_UNCONFIRMED)
        self.assertEqual(result.search_keyword, "广州市白云区点绘化妆品有限公司 阿里巴巴")
        self.assertIn("搜索源出现验证", result.remark)

    def test_search_index_result_marks_exists_on_exact_1688_public_index(self):
        provider = FakeSearchIndexProvider(
            {
                "广州庄雅生物科技有限公司 阿里巴巴": [
                    SearchIndexResult(
                        title="广州庄雅生物科技有限公司-企业信息查询黄页-阿里巴巴",
                        url="https://www.1688.com/factory/example.html",
                        snippet="阿里巴巴提供广州庄雅生物科技有限公司主营沐浴露、洗发水等信息。",
                    )
                ]
            }
        )
        service = Alibaba1688ExistenceService(collector=object(), search_index_provider=provider)

        result = service.verify_company_on_public_search_index("广州庄雅生物科技有限公司")

        self.assertEqual(result.status, ALIBABA_1688_STATUS_EXISTS)
        self.assertEqual(result.matched_company_name, "广州庄雅生物科技有限公司")
        self.assertEqual(result.search_url, "https://www.1688.com/factory/example.html")
        self.assertEqual(result.search_keyword, "广州庄雅生物科技有限公司 阿里巴巴")
        self.assertIn("Bing/微软搜索", result.remark)

    def test_search_index_result_marks_exists_when_search_page_has_any_1688_link(self):
        provider = FakeSearchIndexProvider(
            {
                "广东袋鼠妈妈生物科技有限公司 阿里巴巴": [
                    SearchIndexResult(
                        title="袋鼠 OEM代工厂 - 阿里巴巴",
                        url="https://shop123456.1688.com/",
                        snippet="该搜索结果页出现1688店铺链接，即判定在1688存在。",
                    )
                ]
            }
        )
        service = Alibaba1688ExistenceService(collector=object(), search_index_provider=provider)

        result = service.verify_company_on_public_search_index("广东袋鼠妈妈生物科技有限公司")

        self.assertEqual(result.status, ALIBABA_1688_STATUS_EXISTS)
        self.assertEqual(result.matched_company_name, "广东袋鼠妈妈生物科技有限公司")
        self.assertEqual(result.search_url, "https://shop123456.1688.com/")
        self.assertIn("搜索页出现1688链接", result.remark)

    def test_search_index_result_ignores_generic_1688_homepage(self):
        provider = FakeSearchIndexProvider(
            {
                "广东袋鼠妈妈生物科技有限公司 阿里巴巴": [
                    SearchIndexResult(
                        title="1688.com",
                        url="https://www.1688.com/",
                        snippet="阿里巴巴1688首页。",
                    )
                ]
            }
        )
        service = Alibaba1688ExistenceService(collector=object(), search_index_provider=provider)

        result = service.verify_company_on_public_search_index("广东袋鼠妈妈生物科技有限公司")

        self.assertEqual(result.status, ALIBABA_1688_STATUS_NOT_EXISTS)
        self.assertEqual(result.matched_company_name, "")

    def test_search_index_result_ignores_non_1688_results(self):
        provider = FakeSearchIndexProvider(
            {
                "珠海樱草园生物科技有限公司 阿里巴巴": [
                    SearchIndexResult(
                        title="珠海樱草园生物科技有限公司 - 企查查",
                        url="https://www.qcc.com/firm/example.html",
                        snippet="珠海樱草园生物科技有限公司工商信息。",
                    )
                ]
            }
        )
        service = Alibaba1688ExistenceService(collector=object(), search_index_provider=provider)

        result = service.verify_company_on_public_search_index("珠海樱草园生物科技有限公司")

        self.assertEqual(result.status, ALIBABA_1688_STATUS_NOT_EXISTS)
        self.assertEqual(result.matched_company_name, "")
        self.assertEqual(result.search_keyword, "珠海樱草园生物科技有限公司 阿里巴巴")
        self.assertIn("未发现1688公开索引", result.remark)

    def test_search_index_provider_error_marks_unconfirmed(self):
        provider = FakeSearchIndexProvider(error=RuntimeError("搜索源出现验证或访问限制"))
        service = Alibaba1688ExistenceService(collector=object(), search_index_provider=provider)

        result = service.verify_company_row_with_fallback({"company_name": "广州庄雅生物科技有限公司"})

        self.assertEqual(result.status, ALIBABA_1688_STATUS_UNCONFIRMED)
        self.assertEqual(result.search_keyword, "广州庄雅生物科技有限公司 阿里巴巴")
        self.assertIn("搜索源出现验证", result.remark)

    def test_result_workbook_keeps_source_columns_and_appends_evidence_columns(self):
        service = Alibaba1688ExistenceService(collector=object())
        result = service.judge_company_result(
            "广州市白云区点绘化妆品有限公司",
            ["广州市白云区点绘化妆品有限公司"],
            "https://example.test/search",
            "广州市白云区点绘化妆品有限公司",
        )

        workbook_bytes = service.write_result_workbook(
            ["行业", "公司名称", "备注"],
            [
                {
                    "row_index": 2,
                    "company_name": "广州市白云区点绘化妆品有限公司",
                    "row_values": ["化妆品", "广州市白云区点绘化妆品有限公司", "原表备注"],
                }
            ],
            [result],
        )

        workbook = load_workbook(io.BytesIO(workbook_bytes), data_only=True)
        sheet = workbook["全部结果"]
        header_list = [cell.value for cell in sheet[1]]
        value_list = [cell.value for cell in sheet[2]]

        self.assertEqual(header_list[:3], ["行业", "公司名称", "备注"])
        self.assertEqual(
            header_list[3:],
            [
                "1688反证状态",
                "1688匹配公司名",
                "1688证据链接",
                "1688检索词",
                "1688反证备注",
                "1688反证时间",
            ],
        )
        self.assertEqual(value_list[:3], ["化妆品", "广州市白云区点绘化妆品有限公司", "原表备注"])
        self.assertEqual(value_list[3], ALIBABA_1688_STATUS_EXISTS)
        self.assertEqual(value_list[4], "广州市白云区点绘化妆品有限公司")
        self.assertEqual(value_list[5], "https://example.test/search")
        self.assertEqual(value_list[6], "广州市白云区点绘化妆品有限公司")

    def test_verify_excel_returns_separate_exists_and_not_exists_exports(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["公司名称", "备注"])
        sheet.append(["广州市存在公司有限公司", "应在1688存在"])
        sheet.append(["广州市不存在公司有限公司", "应在1688不存在"])
        excel_output = io.BytesIO()
        workbook.save(excel_output)

        with tempfile.TemporaryDirectory() as temp_dir:
            service = Alibaba1688ExistenceService(
                collector=DeterministicCollector(),
                config=SimpleNamespace(export_dir=Path(temp_dir)),
                search_index_provider=FakeSearchIndexProvider(
                    {
                        "广州市存在公司有限公司 阿里巴巴": [
                            SearchIndexResult(
                                title="广州市存在公司有限公司-企业信息查询黄页-阿里巴巴",
                                url="https://shop.1688.com/example.html",
                                snippet="阿里巴巴收录广州市存在公司有限公司企业信息。",
                            )
                        ],
                        "广州市不存在公司有限公司 阿里巴巴": [
                            SearchIndexResult(
                                title="广州市不存在公司有限公司 - 企查查",
                                url="https://www.qcc.com/firm/example.html",
                                snippet="企业工商信息。",
                            )
                        ],
                    }
                ),
            )

            summary = service.verify_excel_base64(
                base64.b64encode(excel_output.getvalue()).decode("ascii"),
                "反证测试.xlsx",
                company_name_column_index=1,
            )

            self.assertEqual(summary["exists_count"], 1)
            self.assertEqual(summary["not_exists_count"], 1)
            self.assertIn("exists_download_url", summary)
            self.assertIn("not_exists_download_url", summary)
            self.assertTrue(Path(summary["exists_file_path"]).exists())
            self.assertTrue(Path(summary["not_exists_file_path"]).exists())

            exists_workbook = load_workbook(summary["exists_file_path"], data_only=True)
            not_exists_workbook = load_workbook(summary["not_exists_file_path"], data_only=True)
            self.assertEqual(exists_workbook.active.max_row, 2)
            self.assertEqual(not_exists_workbook.active.max_row, 2)
            self.assertEqual(exists_workbook.active["A2"].value, "广州市存在公司有限公司")
            self.assertEqual(not_exists_workbook.active["A2"].value, "广州市不存在公司有限公司")

    def test_verification_job_returns_pending_snapshot_then_completed_results(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["公司名称"])
        sheet.append(["广州市存在公司有限公司"])
        sheet.append(["广州市不存在公司有限公司"])
        excel_output = io.BytesIO()
        workbook.save(excel_output)

        with tempfile.TemporaryDirectory() as temp_dir:
            service = Alibaba1688ExistenceService(
                collector=object(),
                config=SimpleNamespace(export_dir=Path(temp_dir), alibaba_1688_request_interval_seconds=0),
                search_index_provider=FakeSearchIndexProvider(
                    {
                        "广州市存在公司有限公司 阿里巴巴": [
                            SearchIndexResult(
                                title="店铺首页",
                                url="https://example.1688.com/",
                                snippet="1688公开链接",
                            )
                        ],
                        "广州市不存在公司有限公司 阿里巴巴": [
                            SearchIndexResult(
                                title="企业工商信息",
                                url="https://www.qcc.com/firm/example.html",
                                snippet="非1688搜索结果。",
                            )
                        ],
                    }
                ),
            )

            snapshot = service.create_verification_job(
                base64.b64encode(excel_output.getvalue()).decode("ascii"),
                "反证任务.xlsx",
                company_name_column_index=1,
            )

            self.assertEqual(snapshot["status"], "pending")
            self.assertEqual(snapshot["processed_count"], 0)
            self.assertEqual(snapshot["total_count"], 2)
            self.assertEqual(snapshot["result_items"], [])

            service.run_verification_job(snapshot["job_id"])
            completed_snapshot = service.get_verification_job_snapshot(snapshot["job_id"])

            self.assertEqual(completed_snapshot["status"], "completed")
            self.assertEqual(completed_snapshot["processed_count"], 2)
            self.assertEqual(completed_snapshot["exists_count"], 1)
            self.assertEqual(completed_snapshot["not_exists_count"], 1)
            self.assertEqual(len(completed_snapshot["result_items"]), 2)
            self.assertIn("download_url", completed_snapshot)


if __name__ == "__main__":
    unittest.main()
