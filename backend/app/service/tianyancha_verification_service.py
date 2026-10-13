import base64
import binascii
import io
import random
import re
import time
import uuid
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from threading import Event, Lock
from typing import Callable, Dict, List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

from app.core.config import get_config
from app.model.schema import CollectedLeadDTO
from app.service.alibaba_1688_existence_service import find_company_name_column_index
from app.service.collector_service import EnterprisePublicWebCollector, clean_company_name_text


TIANYANCHA_COMPANY_NAME_HEADER_TOKEN_LIST = ["公司", "企业", "客户名称", "单位名称", "名称"]

TIANYANCHA_STATUS_MATCHED = "已匹配"
TIANYANCHA_STATUS_NOT_FOUND = "未找到"
TIANYANCHA_STATUS_UNCONFIRMED = "未确认"

TIANYANCHA_JOB_STATUS_PENDING = "pending"
TIANYANCHA_JOB_STATUS_RUNNING = "running"
TIANYANCHA_JOB_STATUS_WAITING_MANUAL = "waiting_manual"
TIANYANCHA_JOB_STATUS_STOPPING = "stopping"
TIANYANCHA_JOB_STATUS_STOPPED = "stopped"
TIANYANCHA_JOB_STATUS_COMPLETED = "completed"
TIANYANCHA_JOB_STATUS_FAILED = "failed"

TIANYANCHA_RESULT_HEADER_LIST = [
    "天眼查反证状态",
    "天眼查匹配企业",
    "法定代表人",
    "联系人",
    "联系电话",
    "电话来源",
    "企业状态",
    "统一社会信用代码",
    "企业地址",
    "经营范围",
    "风险标签",
    "天眼查详情页",
    "天眼查检索词",
    "天眼查反证备注",
    "天眼查反证时间",
]

TIANYANCHA_EXPORT_HEADER_LIST = [
    "企业名称",
    "天眼查反证状态",
    "天眼查匹配企业",
    "法定代表人",
    "联系人",
    "最大股东",
    "联系电话",
    "电话来源",
    "企业地址",
    "经营范围",
    "天眼查反证时间",
]

TIANYANCHA_EXPORT_TABLE_NAME_MAP = {
    "全部结果": "TianyanchaAllResultTable",
    TIANYANCHA_STATUS_MATCHED: "TianyanchaMatchedTable",
    TIANYANCHA_STATUS_NOT_FOUND: "TianyanchaNotFoundTable",
    TIANYANCHA_STATUS_UNCONFIRMED: "TianyanchaUnconfirmedTable",
}
TIANYANCHA_CONTACT_ROLE_PATTERN = re.compile(
    r"^\s*(法人/最大股东|法人|最大股东)\s*[：:]\s*(.*?)\s*$"
)

TIANYANCHA_RESULT_SHEET_NAME_LIST = [
    "全部结果",
    TIANYANCHA_STATUS_MATCHED,
    TIANYANCHA_STATUS_NOT_FOUND,
    TIANYANCHA_STATUS_UNCONFIRMED,
]


@dataclass
class TianyanchaVerificationResult:
    source_company_name: str
    status: str
    matched_company_name: str = ""
    legal_person: str = ""
    contact_name: str = ""
    max_shareholder: str = ""
    phone: str = ""
    phone_source: str = ""
    company_status: str = ""
    unified_social_credit_code: str = ""
    registered_address: str = ""
    business_scope: str = ""
    risk_tag: str = ""
    profile_url: str = ""
    search_keyword: str = ""
    remark: str = ""
    verified_at: str = ""


@dataclass
class TianyanchaVerificationJob:
    job_id: str
    file_name: str
    header_list: List
    company_row_list: List[Dict]
    result_list: List[TianyanchaVerificationResult]
    batch_size: int
    jitter_min_seconds: float
    jitter_max_seconds: float
    batch_cooldown_seconds: float
    status: str = TIANYANCHA_JOB_STATUS_PENDING
    current_row_index: int = 0
    current_company_name: str = ""
    current_search_keyword: str = ""
    current_batch_index: int = 0
    total_batch_count: int = 0
    current_message: str = ""
    error_message: str = ""
    file_path: str = ""
    file_name_export: str = ""
    download_url: str = ""
    created_at: str = ""
    updated_at: str = ""
    event_list: List[Dict] = field(default_factory=list)
    stop_requested: bool = False
    stop_event: Event = field(default_factory=Event, repr=False)


def find_tianyancha_company_name_column_index(header_list: List[str], manual_column_index: int = 0) -> int:
    detected_index = find_company_name_column_index(header_list, manual_column_index)
    if detected_index > 0:
        return detected_index
    normalized_header_list = [str(header or "").strip() for header in header_list or []]
    for index, header in enumerate(normalized_header_list, start=1):
        if not header:
            continue
        if any(token in header for token in TIANYANCHA_COMPANY_NAME_HEADER_TOKEN_LIST):
            return index
    return 0


class TianyanchaVerificationService:
    def __init__(
        self,
        collector_factory: Optional[Callable[[], EnterprisePublicWebCollector]] = None,
        config=None,
        sleep_func: Callable[[float], None] = time.sleep,
        random_uniform_func: Callable[[float, float], float] = random.uniform,
    ):
        self.collector_factory = collector_factory or EnterprisePublicWebCollector.create_tianyancha_collector
        self.config = config or get_config()
        self.sleep_func = sleep_func
        self.random_uniform_func = random_uniform_func
        self.job_lock = Lock()
        self.job_map: Dict[str, TianyanchaVerificationJob] = {}

    def preview_excel_columns_base64(self, excel_base64: str, file_name: str) -> Dict:
        excel_bytes = self.decode_excel_base64(excel_base64, file_name)
        workbook = load_workbook(io.BytesIO(excel_bytes), read_only=True, data_only=True)
        sheet = workbook.active
        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), [])
        header_list = ["" if value is None else str(value).strip() for value in header_row]
        suggested_column_index = find_tianyancha_company_name_column_index(header_list)
        suggested_column_name = (
            header_list[suggested_column_index - 1]
            if 0 < suggested_column_index <= len(header_list)
            else ""
        )
        return {
            "headers": header_list,
            "suggested_column_index": suggested_column_index,
            "suggested_column_name": suggested_column_name,
        }

    def create_verification_job(
        self,
        excel_base64: str,
        file_name: str,
        company_name_column_index: int = 0,
        batch_size: int = 3,
        jitter_min_seconds: float = 8.0,
        jitter_max_seconds: float = 22.0,
        batch_cooldown_seconds: float = 45.0,
    ) -> Dict:
        excel_bytes = self.decode_excel_base64(excel_base64, file_name)
        excel_data = self.read_company_rows_from_excel_bytes(
            excel_bytes,
            company_name_column_index=company_name_column_index,
        )
        clean_batch_size = max(1, min(int(batch_size or 3), 30))
        clean_jitter_min = max(0.0, min(float(jitter_min_seconds or 0.0), 900.0))
        clean_jitter_max = max(0.0, min(float(jitter_max_seconds or 0.0), 1200.0))
        if clean_jitter_max < clean_jitter_min:
            clean_jitter_min, clean_jitter_max = clean_jitter_max, clean_jitter_min
        clean_cooldown = max(0.0, min(float(batch_cooldown_seconds or 0.0), 1800.0))
        total_count = len(excel_data["company_row_list"])
        created_at = self.build_verified_at()
        job = TianyanchaVerificationJob(
            job_id=uuid.uuid4().hex,
            file_name=file_name,
            header_list=excel_data["header_list"],
            company_row_list=excel_data["company_row_list"],
            result_list=[],
            batch_size=clean_batch_size,
            jitter_min_seconds=clean_jitter_min,
            jitter_max_seconds=clean_jitter_max,
            batch_cooldown_seconds=clean_cooldown,
            total_batch_count=(total_count + clean_batch_size - 1) // clean_batch_size if total_count else 0,
            created_at=created_at,
            updated_at=created_at,
        )
        self.append_job_event(job, "info", f"天眼查反证任务已创建：{total_count} 家企业")
        with self.job_lock:
            self.job_map[job.job_id] = job
            return self.build_job_snapshot(job)

    def run_verification_job(self, job_id: str) -> None:
        with self.job_lock:
            job = self.get_verification_job(job_id)
            if job.stop_requested:
                job.status = TIANYANCHA_JOB_STATUS_STOPPING
                job.current_message = "正在安全停止，并保存已完成结果"
            else:
                job.status = TIANYANCHA_JOB_STATUS_RUNNING
                job.current_message = "正在启动天眼查反证"
            job.error_message = ""
        try:
            if self.is_stop_requested(job):
                self.finish_stopped_job(job)
                return
            if not job.company_row_list:
                self.finish_job(job)
                return
            for batch_start in range(0, len(job.company_row_list), job.batch_size):
                if self.is_stop_requested(job):
                    self.finish_stopped_job(job)
                    return
                batch_index = batch_start // job.batch_size + 1
                batch_row_list = job.company_row_list[batch_start : batch_start + job.batch_size]
                self.mark_batch_start(job, batch_index, batch_row_list)
                try:
                    batch_lead_list = self.collect_batch(job, batch_row_list)
                    if self.is_stop_requested(job):
                        self.finish_stopped_job(job)
                        return
                    batch_result_map = self.build_result_map(batch_lead_list)
                    for row_offset, company_row in enumerate(batch_row_list):
                        if self.is_stop_requested(job):
                            self.finish_stopped_job(job)
                            return
                        company_name = company_row.get("company_name", "")
                        with self.job_lock:
                            job.current_row_index = company_row.get("row_index", batch_start + row_offset + 2)
                            job.current_company_name = company_name
                            job.current_search_keyword = company_name
                        lead = self.find_result_lead(company_name, batch_result_map)
                        result = self.build_result_from_lead(company_name, lead)
                        with self.job_lock:
                            job.result_list.append(result)
                except Exception as ex:
                    if self.is_stop_requested(job):
                        self.finish_stopped_job(job)
                        return
                    error_message = str(ex)[:220]
                    self.append_job_event(job, "warning", f"第 {batch_index} 批未完成：{error_message}")
                    with self.job_lock:
                        for company_row in batch_row_list:
                            job.result_list.append(
                                self.build_unconfirmed_result(
                                    company_row.get("company_name", ""),
                                    error_message,
                                )
                            )
                if batch_start + job.batch_size < len(job.company_row_list):
                    self.wait_before_next_batch(job, batch_index + 1)
                    if self.is_stop_requested(job):
                        self.finish_stopped_job(job)
                        return
            self.finish_job(job)
        except Exception as ex:
            if self.is_stop_requested(job):
                self.finish_stopped_job(job)
                return
            with self.job_lock:
                job.status = TIANYANCHA_JOB_STATUS_FAILED
                job.error_message = str(ex)[:300]
                job.current_message = "天眼查反证失败"

    def has_active_job(self) -> bool:
        with self.job_lock:
            return any(
                job.status
                in {
                    TIANYANCHA_JOB_STATUS_PENDING,
                    TIANYANCHA_JOB_STATUS_RUNNING,
                    TIANYANCHA_JOB_STATUS_WAITING_MANUAL,
                    TIANYANCHA_JOB_STATUS_STOPPING,
                }
                for job in self.job_map.values()
            )

    def collect_batch(self, job: TianyanchaVerificationJob, company_row_list: List[Dict]) -> List[CollectedLeadDTO]:
        collector = self.collector_factory()
        set_callback = getattr(collector, "set_task_event_callback", None)
        if callable(set_callback):
            set_callback(lambda event_type, level, message: self.handle_collector_event(job, event_type, level, message))
        set_stop_callback = getattr(collector, "set_stop_requested_callback", None)
        if callable(set_stop_callback):
            set_stop_callback(lambda: self.is_stop_requested(job))
        company_name_list = [row.get("company_name", "") for row in company_row_list]
        return collector.collect_by_company_names(
            industry="",
            company_name_list=company_name_list,
            collect_limit=len(company_name_list),
        )

    def mark_batch_start(self, job: TianyanchaVerificationJob, batch_index: int, batch_row_list: List[Dict]) -> None:
        first_row = batch_row_list[0] if batch_row_list else {}
        first_company_name = first_row.get("company_name", "")
        with self.job_lock:
            if job.stop_requested:
                return
            job.status = TIANYANCHA_JOB_STATUS_RUNNING
            job.current_batch_index = batch_index
            job.current_row_index = first_row.get("row_index", 0)
            job.current_company_name = first_company_name
            job.current_search_keyword = first_company_name
            job.current_message = f"正在处理第 {batch_index}/{job.total_batch_count} 批"
        self.append_job_event(job, "info", f"开始第 {batch_index}/{job.total_batch_count} 批：{len(batch_row_list)} 家企业")

    def wait_before_next_batch(self, job: TianyanchaVerificationJob, next_batch_index: int) -> None:
        jitter_seconds = self.random_uniform_func(job.jitter_min_seconds, job.jitter_max_seconds)
        wait_seconds = round(job.batch_cooldown_seconds + jitter_seconds, 1)
        if wait_seconds <= 0:
            return
        with self.job_lock:
            if job.stop_requested:
                return
            job.status = TIANYANCHA_JOB_STATUS_RUNNING
            job.current_message = f"批次冷却中，等待 {wait_seconds} 秒后进入第 {next_batch_index} 批"
        self.append_job_event(job, "info", f"批次冷却 {wait_seconds} 秒，降低连续访问频率")
        if self.sleep_func is time.sleep:
            job.stop_event.wait(wait_seconds)
        else:
            self.sleep_func(wait_seconds)

    def handle_collector_event(
        self,
        job: TianyanchaVerificationJob,
        event_type: str,
        level: str,
        message: str,
    ) -> None:
        clean_message = str(message or "").strip()
        with self.job_lock:
            if job.stop_requested:
                job.status = TIANYANCHA_JOB_STATUS_STOPPING
            elif event_type == "waiting_manual":
                job.status = TIANYANCHA_JOB_STATUS_WAITING_MANUAL
            elif event_type == "manual_resumed":
                job.status = TIANYANCHA_JOB_STATUS_RUNNING
            job.current_message = clean_message[:260]
        self.append_job_event(job, level or "info", clean_message)

    def finish_job(self, job: TianyanchaVerificationJob) -> None:
        if self.is_stop_requested(job):
            self.finish_stopped_job(job)
            return
        summary = self.write_result_exports(job.header_list, job.company_row_list, job.result_list, job.file_name)
        with self.job_lock:
            job.status = TIANYANCHA_JOB_STATUS_COMPLETED
            job.current_row_index = 0
            job.current_company_name = ""
            job.current_search_keyword = ""
            job.current_message = "天眼查反证完成"
            job.file_path = summary.get("file_path", "")
            job.file_name_export = summary.get("file_name", "")
            job.download_url = summary.get("download_url", "")
        self.append_job_event(job, "info", "天眼查反证完成，已生成结果表")

    def stop_verification_job(self, job_id: str) -> Dict:
        with self.job_lock:
            job = self.get_verification_job(job_id)
            if job.status in {
                TIANYANCHA_JOB_STATUS_COMPLETED,
                TIANYANCHA_JOB_STATUS_FAILED,
                TIANYANCHA_JOB_STATUS_STOPPED,
            }:
                return self.build_job_snapshot(job)
            job.stop_requested = True
            job.status = TIANYANCHA_JOB_STATUS_STOPPING
            job.current_message = "正在安全停止，并保存已完成结果"
            job.stop_event.set()
        self.append_job_event(job, "warning", "收到停止请求，正在保存已完成的反证结果")
        return self.get_verification_job_snapshot(job_id)

    def is_stop_requested(self, job: TianyanchaVerificationJob) -> bool:
        return job.stop_requested or job.stop_event.is_set()

    def finish_stopped_job(self, job: TianyanchaVerificationJob) -> None:
        processed_count = len(job.result_list)
        summary = self.write_result_exports(
            job.header_list,
            job.company_row_list[:processed_count],
            job.result_list,
            job.file_name,
        )
        with self.job_lock:
            job.status = TIANYANCHA_JOB_STATUS_STOPPED
            job.current_row_index = 0
            job.current_company_name = ""
            job.current_search_keyword = ""
            job.current_message = f"已安全停止，已保存 {processed_count} 条结果"
            job.file_path = summary.get("file_path", "")
            job.file_name_export = summary.get("file_name", "")
            job.download_url = summary.get("download_url", "")
        self.append_job_event(job, "info", f"天眼查反证已停止，已生成 {processed_count} 条部分结果")

    def get_verification_job_snapshot(self, job_id: str) -> Dict:
        with self.job_lock:
            return self.build_job_snapshot(self.get_verification_job(job_id))

    def get_latest_verification_job_snapshot(self) -> Dict:
        with self.job_lock:
            if not self.job_map:
                return self.build_empty_job_snapshot()
            latest_job = next(reversed(self.job_map.values()))
            return self.build_job_snapshot(latest_job)

    def list_verification_job_history(self, limit: int = 20) -> List[Dict]:
        clean_limit = max(1, min(int(limit or 20), 100))
        with self.job_lock:
            job_list = list(self.job_map.values())[-clean_limit:]
            return [self.build_job_history_item(job) for job in reversed(job_list)]

    def get_verification_job(self, job_id: str) -> TianyanchaVerificationJob:
        job = self.job_map.get(job_id)
        if job is None:
            raise KeyError(job_id)
        return job

    def decode_excel_base64(self, excel_base64: str, file_name: str) -> bytes:
        if not str(file_name or "").lower().endswith(".xlsx"):
            raise ValueError("仅支持 .xlsx 文件")
        try:
            return base64.b64decode(excel_base64, validate=True)
        except (binascii.Error, ValueError) as ex:
            raise ValueError("Excel base64 内容无效") from ex

    def read_company_rows_from_excel_bytes(self, excel_bytes: bytes, company_name_column_index: int = 0) -> Dict:
        workbook = load_workbook(io.BytesIO(excel_bytes), read_only=True, data_only=True)
        sheet = workbook.active
        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), [])
        header_list = ["" if value is None else str(value).strip() for value in header_row]
        detected_column_index = find_tianyancha_company_name_column_index(header_list, company_name_column_index)
        if detected_column_index <= 0:
            raise ValueError("请选择公司名称列")

        company_row_list = []
        for row_index in range(2, sheet.max_row + 1):
            row_values = [
                sheet.cell(row=row_index, column=column_index).value
                for column_index in range(1, len(header_list) + 1)
            ]
            company_name = row_values[detected_column_index - 1]
            clean_company_name = clean_company_name_text("" if company_name is None else str(company_name).strip())
            if not clean_company_name:
                continue
            company_row_list.append(
                {
                    "row_index": row_index,
                    "company_name": clean_company_name,
                    "row_values": row_values,
                }
            )
        return {
            "header_list": header_list,
            "company_name_column_index": detected_column_index,
            "company_row_list": company_row_list,
        }

    def build_result_map(self, lead_list: List[CollectedLeadDTO]) -> Dict[str, CollectedLeadDTO]:
        result = {}
        for lead in lead_list or []:
            key = EnterprisePublicWebCollector.normalize_company_match_key(lead.suspected_company_name)
            if key and key not in result:
                result[key] = lead
        return result

    def find_result_lead(self, company_name: str, result_map: Dict[str, CollectedLeadDTO]) -> Optional[CollectedLeadDTO]:
        key = EnterprisePublicWebCollector.normalize_company_match_key(company_name)
        return result_map.get(key)

    def build_result_from_lead(
        self,
        company_name: str,
        lead: Optional[CollectedLeadDTO],
    ) -> TianyanchaVerificationResult:
        if lead is None:
            return TianyanchaVerificationResult(
                source_company_name=company_name,
                status=TIANYANCHA_STATUS_NOT_FOUND,
                search_keyword=company_name,
                remark="未找到精准匹配企业，或页面没有可见企业信息",
                verified_at=self.build_verified_at(),
            )
        return TianyanchaVerificationResult(
            source_company_name=company_name,
            status=TIANYANCHA_STATUS_MATCHED,
            matched_company_name=lead.suspected_company_name,
            legal_person=lead.legal_person,
            contact_name=lead.contact_name,
            max_shareholder=self.extract_export_contact_fields(lead.contact_name, lead.legal_person)[1],
            phone=lead.phone,
            phone_source=lead.phone_source,
            company_status=lead.company_status,
            unified_social_credit_code=lead.unified_social_credit_code,
            registered_address=lead.registered_address,
            business_scope=lead.business_scope,
            risk_tag=lead.risk_tag,
            profile_url=lead.profile_url,
            search_keyword=lead.source_keyword or company_name,
            remark="天眼查公开页精准匹配成功",
            verified_at=self.build_verified_at(),
        )

    def build_unconfirmed_result(self, company_name: str, remark: str) -> TianyanchaVerificationResult:
        return TianyanchaVerificationResult(
            source_company_name=company_name,
            status=TIANYANCHA_STATUS_UNCONFIRMED,
            search_keyword=company_name,
            remark=remark,
            verified_at=self.build_verified_at(),
        )

    def write_result_exports(
        self,
        header_list: List,
        company_row_list: List[Dict],
        result_list: List[TianyanchaVerificationResult],
        file_name: str,
    ) -> Dict:
        result_workbook_bytes = self.write_result_workbook(header_list, company_row_list, result_list)
        export_file_path = self.build_export_file_path(file_name)
        export_file_path.write_bytes(result_workbook_bytes)
        return self.build_summary(company_row_list, result_list, export_file_path)

    def write_result_workbook(
        self,
        header_list: List,
        company_row_list: List[Dict],
        result_list: List[TianyanchaVerificationResult],
    ) -> bytes:
        workbook = Workbook()
        workbook.active.title = TIANYANCHA_RESULT_SHEET_NAME_LIST[0]
        for sheet_name in TIANYANCHA_RESULT_SHEET_NAME_LIST[1:]:
            workbook.create_sheet(sheet_name)
        row_result_pair_list = list(zip(company_row_list, result_list))
        self.write_result_sheet(workbook[TIANYANCHA_RESULT_SHEET_NAME_LIST[0]], header_list, row_result_pair_list)
        for status in (
            TIANYANCHA_STATUS_MATCHED,
            TIANYANCHA_STATUS_NOT_FOUND,
            TIANYANCHA_STATUS_UNCONFIRMED,
        ):
            self.write_result_sheet(
                workbook[status],
                header_list,
                [(row, result) for row, result in row_result_pair_list if result.status == status],
            )
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()

    def write_result_sheet(self, sheet, header_list: List, row_result_pair_list: List) -> None:
        del header_list
        sheet.append(TIANYANCHA_EXPORT_HEADER_LIST)
        for company_row, result in row_result_pair_list:
            contact_name, max_shareholder = self.extract_export_contact_fields(
                result.contact_name,
                result.legal_person,
            )
            sheet.append(
                [
                    result.source_company_name or company_row.get("company_name", ""),
                    result.status,
                    result.matched_company_name,
                    result.legal_person,
                    contact_name,
                    result.max_shareholder or max_shareholder,
                    result.phone,
                    result.phone_source,
                    result.registered_address,
                    result.business_scope,
                    result.verified_at,
                ]
            )
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:K{max(1, sheet.max_row)}"
        for column_letter, width in zip(
            "ABCDEFGHIJK",
            (28, 16, 28, 14, 18, 18, 36, 32, 42, 60, 22),
        ):
            sheet.column_dimensions[column_letter].width = width
        table_name = TIANYANCHA_EXPORT_TABLE_NAME_MAP.get(sheet.title)
        if table_name and sheet.max_row >= 2:
            table = Table(displayName=table_name, ref=f"A1:K{sheet.max_row}")
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            sheet.add_table(table)

    def extract_export_contact_fields(self, contact_name: str, legal_person: str) -> tuple[str, str]:
        clean_contact_name = str(contact_name or "").strip()
        clean_legal_person = str(legal_person or "").strip()
        if not clean_contact_name:
            return clean_legal_person, ""
        contact_person = ""
        max_shareholder = ""
        has_role_label = False
        for part in re.split(r"[；;]", clean_contact_name):
            match = TIANYANCHA_CONTACT_ROLE_PATTERN.fullmatch(part.strip())
            if not match:
                continue
            has_role_label = True
            role_name, person_name = match.groups()
            person_name = person_name.strip()
            if not person_name or person_name in {"万元", "万"}:
                continue
            if role_name in {"法人", "法人/最大股东"}:
                contact_person = person_name
            if role_name in {"最大股东", "法人/最大股东"}:
                max_shareholder = person_name
        if has_role_label:
            return contact_person or clean_legal_person, max_shareholder
        return clean_contact_name, ""

    def build_export_file_path(self, file_name: str) -> Path:
        export_dir = Path(self.config.export_dir)
        export_dir.mkdir(parents=True, exist_ok=True)
        stem = Path(file_name).stem or "tianyancha_verification"
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        return export_dir / f"{stem}_tianyancha_verification_{timestamp}.xlsx"

    def build_summary(
        self,
        company_row_list: List[Dict],
        result_list: List[TianyanchaVerificationResult],
        export_file_path: Path,
    ) -> Dict:
        return {
            "total_count": len(company_row_list),
            "matched_count": sum(1 for result in result_list if result.status == TIANYANCHA_STATUS_MATCHED),
            "not_found_count": sum(1 for result in result_list if result.status == TIANYANCHA_STATUS_NOT_FOUND),
            "unconfirmed_count": sum(1 for result in result_list if result.status == TIANYANCHA_STATUS_UNCONFIRMED),
            "result_items": self.build_result_item_list(company_row_list, result_list),
            "file_path": str(export_file_path),
            "file_name": export_file_path.name,
            "download_url": f"/api/exports/{export_file_path.name}",
        }

    def build_job_snapshot(self, job: TianyanchaVerificationJob) -> Dict:
        result_list = list(job.result_list)
        processed_count = len(result_list)
        return {
            "job_id": job.job_id,
            "status": job.status,
            "processed_count": processed_count,
            "current_row_index": job.current_row_index,
            "current_company_name": job.current_company_name,
            "current_search_keyword": job.current_search_keyword,
            "current_batch_index": job.current_batch_index,
            "total_batch_count": job.total_batch_count,
            "current_message": job.current_message,
            "error_message": job.error_message,
            "file_name": Path(job.file_name).name,
            "file_path": job.file_path,
            "download_url": job.download_url,
            "export_file_name": job.file_name_export,
            "created_at": job.created_at,
            "updated_at": job.updated_at,
            "total_count": len(job.company_row_list),
            "matched_count": sum(1 for result in result_list if result.status == TIANYANCHA_STATUS_MATCHED),
            "not_found_count": sum(1 for result in result_list if result.status == TIANYANCHA_STATUS_NOT_FOUND),
            "unconfirmed_count": sum(1 for result in result_list if result.status == TIANYANCHA_STATUS_UNCONFIRMED),
            "batch_size": job.batch_size,
            "jitter_min_seconds": job.jitter_min_seconds,
            "jitter_max_seconds": job.jitter_max_seconds,
            "batch_cooldown_seconds": job.batch_cooldown_seconds,
            "result_items": self.build_result_item_list(job.company_row_list[:processed_count], result_list),
            "event_items": list(job.event_list[-80:]),
        }

    def build_job_history_item(self, job: TianyanchaVerificationJob) -> Dict:
        result_list = list(job.result_list)
        return {
            "job_id": job.job_id,
            "status": job.status,
            "file_name": Path(job.file_name).name,
            "download_url": job.download_url,
            "export_file_name": job.file_name_export,
            "total_count": len(job.company_row_list),
            "processed_count": len(result_list),
            "matched_count": sum(1 for result in result_list if result.status == TIANYANCHA_STATUS_MATCHED),
            "not_found_count": sum(1 for result in result_list if result.status == TIANYANCHA_STATUS_NOT_FOUND),
            "unconfirmed_count": sum(1 for result in result_list if result.status == TIANYANCHA_STATUS_UNCONFIRMED),
            "current_message": job.current_message,
            "error_message": job.error_message,
            "created_at": job.created_at,
            "updated_at": job.updated_at,
        }

    def build_empty_job_snapshot(self) -> Dict:
        return {
            "job_id": "",
            "status": "none",
            "processed_count": 0,
            "current_row_index": 0,
            "current_company_name": "",
            "current_search_keyword": "",
            "current_batch_index": 0,
            "total_batch_count": 0,
            "current_message": "",
            "error_message": "",
            "file_name": "",
            "file_path": "",
            "download_url": "",
            "export_file_name": "",
            "created_at": "",
            "updated_at": "",
            "total_count": 0,
            "matched_count": 0,
            "not_found_count": 0,
            "unconfirmed_count": 0,
            "batch_size": 3,
            "jitter_min_seconds": 8.0,
            "jitter_max_seconds": 22.0,
            "batch_cooldown_seconds": 45.0,
            "result_items": [],
            "event_items": [],
        }

    def build_result_item_list(
        self,
        company_row_list: List[Dict],
        result_list: List[TianyanchaVerificationResult],
    ) -> List[Dict]:
        item_list = []
        for company_row, result in zip(company_row_list, result_list):
            item_list.append(
                {
                    "row_index": company_row.get("row_index", 0),
                    "company_name": company_row.get("company_name", ""),
                    "status": result.status,
                    "matched_company_name": result.matched_company_name,
                    "legal_person": result.legal_person,
                    "contact_name": result.contact_name,
                    "max_shareholder": result.max_shareholder,
                    "phone": result.phone,
                    "phone_source": result.phone_source,
                    "company_status": result.company_status,
                    "unified_social_credit_code": result.unified_social_credit_code,
                    "registered_address": result.registered_address,
                    "business_scope": result.business_scope,
                    "risk_tag": result.risk_tag,
                    "profile_url": result.profile_url,
                    "search_keyword": result.search_keyword,
                    "remark": result.remark,
                    "verified_at": result.verified_at,
                }
            )
        return item_list

    def append_job_event(self, job: TianyanchaVerificationJob, level: str, message: str) -> None:
        clean_message = str(message or "").strip()
        if not clean_message:
            return
        with self.job_lock:
            created_at = self.build_verified_at()
            job.event_list.append(
                {
                    "level": level or "info",
                    "message": clean_message[:300],
                    "created_at": created_at,
                }
            )
            job.updated_at = created_at

    def build_verified_at(self) -> str:
        return datetime.now().strftime("%Y-%m-%d %H:%M:%S")
