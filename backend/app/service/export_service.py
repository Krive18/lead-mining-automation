from pathlib import Path
from typing import Dict, List
from urllib.parse import quote, urlparse

from app.core.config import get_config
from app.core.constant import (
    ASSIGNMENT_STATUS_ASSIGNED,
    SOURCE_TYPE_AI_WEB_SEARCH,
    SOURCE_TYPE_ALIBABA_1688,
    SOURCE_TYPE_ALIBABA_1688_PAGE_ASSISTANT,
    SOURCE_TYPE_ALIBABA_1688_PAGE_IMPORT,
    SOURCE_TYPE_MADE_IN_CHINA_TIANYANCHA,
    SOURCE_TYPE_QIXIN_PUBLIC_WEB,
    SOURCE_TYPE_TIANYANCHA_PUBLIC_WEB,
)
from app.model.schema import ExportCreateDTO
from app.repository.lead_repository import LeadRepository
from app.util.address_util import clean_registered_address_text
from app.util.mask_util import mask_phone
from app.util.time_util import get_file_time_text


EXPORT_HEADER = [
    ("level", "线索等级"),
    ("score", "线索评分"),
    ("assignment_status", "分发状态"),
    ("industry", "行业"),
    ("province", "省份"),
    ("city", "城市"),
    ("district", "区县"),
    ("registered_address", "企业地址"),
    ("city_source", "城市来源"),
    ("company_name", "公司名称"),
    ("suspected_company_name", "疑似公司名"),
    ("product_name", "产品/主营"),
    ("contact_name", "联系人"),
    ("phone", "联系电话"),
    ("phone_source", "电话来源"),
    ("website_url", "官网"),
    ("shop_url", "店铺入口"),
    ("profile_url", "来源主页"),
    ("source_type", "来源平台"),
    ("source_keyword", "来源关键词"),
    ("company_status", "企业状态"),
    ("unified_social_credit_code", "统一社会信用代码"),
    ("legal_person", "法定代表人"),
    ("business_scope", "经营范围"),
    ("risk_tag", "风险标签"),
    ("verification_status", "核验状态"),
    ("confidence", "置信度"),
    ("gmt_create", "采集时间"),
]

_company_name_index = next(index for index, (field_name, _) in enumerate(EXPORT_HEADER) if field_name == "company_name")
EXPORT_HEADER.insert(_company_name_index + 1, ("company_name_en", "英文公司名称"))
_product_name_index = next(index for index, (field_name, _) in enumerate(EXPORT_HEADER) if field_name == "product_name")
EXPORT_HEADER.insert(_product_name_index + 1, ("shop_level", "店铺等级"))

TIANYANCHA_EXPORT_SOURCE_TYPE_SET = {
    SOURCE_TYPE_ALIBABA_1688,
    SOURCE_TYPE_MADE_IN_CHINA_TIANYANCHA,
    SOURCE_TYPE_TIANYANCHA_PUBLIC_WEB,
    SOURCE_TYPE_AI_WEB_SEARCH,
}

ALIBABA_1688_EXPORT_SOURCE_TYPE_SET = {
    SOURCE_TYPE_ALIBABA_1688,
    SOURCE_TYPE_ALIBABA_1688_PAGE_ASSISTANT,
    SOURCE_TYPE_ALIBABA_1688_PAGE_IMPORT,
}

EXPORT_SOURCE_GROUP_LABEL_MAP = {
    "": "全部路线",
    "all": "全部路线",
    "qcc": "企查查",
    "aiqicha": "爱企查",
    "qixin": "启信宝",
    "tianyancha": "天眼查",
    "ai_web": "AI 1688搜企",
    "fengniao": "蜂鸟",
    "douyin": "抖音线索",
    "alibaba_international": "阿里国际站",
    "alibaba_tianyancha": "1688采集",
    "made_in_china": "中国制造网+天眼查",
    "alibaba": "1688普通页面",
}

EXPORT_SOURCE_TYPE_LABEL_MAP = {
    "qcc_public_web": "企查查",
    "aiqicha_public_web": "爱企查",
    "qixin_public_web": "启信宝",
    "tianyancha_public_web": "天眼查",
    "ai_web_search": "AI 1688搜企",
    "fengniao_public_web": "蜂鸟",
    "douyin_qcc": "抖音线索",
    "alibaba_international": "阿里国际站",
    "alibaba_1688": "1688采集",
    "made_in_china_tianyancha": "中国制造网+天眼查",
    "alibaba_1688_page_assistant": "1688普通页面",
    "alibaba_1688_page_import": "1688普通页面",
}


class ExportService:
    def __init__(self, repository: LeadRepository):
        self.repository = repository

    def create_export(self, request: ExportCreateDTO) -> Dict:
        config = get_config()
        levels = request.levels if request.levels else ["A", "B"]
        lead_list = self.list_export_lead(
            levels=levels,
            source_type=request.source_type,
            source_group=request.source_group,
            industry=request.industry,
            keyword=request.keyword,
            city=request.city,
            assignment_status=request.assignment_status,
        )
        export_lead_list = []
        for lead in lead_list:
            export_lead = dict(lead)
            export_lead["assignment_status"] = ASSIGNMENT_STATUS_ASSIGNED
            export_lead_list.append(export_lead)
        route_label = self.build_export_route_label(request.source_group, request.source_type)
        file_path = self.build_unique_export_file_path(config.export_dir, route_label)
        file_name = file_path.name
        self.write_workbook(file_path, export_lead_list, request.is_phone_masked)
        self.repository.update_assignment_status(
            [lead.get("id", "") for lead in lead_list],
            ASSIGNMENT_STATUS_ASSIGNED,
        )
        return {
            "file_name": file_name,
            "file_path": str(file_path),
            "download_url": f"/api/exports/{file_name}",
            "total_count": len(lead_list),
        }

    def build_export_route_label(self, source_group: str = "", source_type: str = "") -> str:
        if source_type:
            return EXPORT_SOURCE_TYPE_LABEL_MAP.get(source_type, "指定路线")
        return EXPORT_SOURCE_GROUP_LABEL_MAP.get(source_group, "全部路线")

    def build_unique_export_file_path(self, export_dir: Path, route_label: str = "全部路线") -> Path:
        export_dir.mkdir(parents=True, exist_ok=True)
        timestamp = get_file_time_text()
        file_path = export_dir / f"{route_label}_线索导出_{timestamp}.xlsx"
        if not file_path.exists():
            return file_path
        for index in range(1, 1000):
            candidate_file_path = export_dir / f"{route_label}_线索导出_{timestamp}_{index}.xlsx"
            if not candidate_file_path.exists():
                return candidate_file_path
        raise RuntimeError("导出文件名冲突过多，请稍后再试")

    def list_export_lead(
        self,
        levels: List[str],
        source_type: str = "",
        source_group: str = "",
        industry: str = "",
        keyword: str = "",
        city: str = "",
        assignment_status: str = "",
    ) -> List[Dict]:
        lead_list = self.repository.list_lead_for_export(
            levels=levels,
            source_type=source_type,
            source_group=source_group,
            industry=industry,
            keyword=keyword,
            city=city,
            assignment_status=assignment_status,
        )
        lead_id_set = set()
        result = []
        for lead in lead_list:
            if lead["id"] in lead_id_set:
                continue
            lead_id_set.add(lead["id"])
            result.append(lead)
        return result

    def format_export_phone_source(self, lead: dict) -> object:
        phone_source = lead.get("phone_source")
        phone_source_text = "" if phone_source is None else str(phone_source).strip()
        if not phone_source_text:
            return ""

        source_type = str(lead.get("source_type") or "").strip()
        if source_type in TIANYANCHA_EXPORT_SOURCE_TYPE_SET:
            return "天眼查"
        if source_type == SOURCE_TYPE_QIXIN_PUBLIC_WEB:
            return "启信宝"
        return phone_source

    def format_export_url(self, lead: dict, field_name: str) -> str:
        source_type = str(lead.get("source_type") or "").strip()
        if source_type not in ALIBABA_1688_EXPORT_SOURCE_TYPE_SET:
            return str(lead.get(field_name) or "").strip()

        shop_url = self.pick_alibaba_1688_url(lead.get("shop_url"))
        profile_url = self.pick_alibaba_1688_url(lead.get("profile_url"))
        if field_name == "shop_url":
            return shop_url
        return profile_url or shop_url or self.build_alibaba_1688_search_url(lead.get("source_keyword"))

    @staticmethod
    def pick_alibaba_1688_url(value) -> str:
        url = str(value or "").strip()
        if not url:
            return ""
        try:
            parsed_url = urlparse(url)
        except ValueError:
            return ""
        host = parsed_url.netloc.lower()
        if host == "1688.com" or host.endswith(".1688.com"):
            return url
        return ""

    @staticmethod
    def build_alibaba_1688_search_url(source_keyword) -> str:
        keyword = str(source_keyword or "").strip()
        if not keyword:
            return ""
        encoded_keyword = quote(keyword, safe="", encoding="gbk", errors="ignore")
        encoded_province = quote("\u5e7f\u4e1c", safe="", encoding="gbk", errors="ignore")
        return (
            "https://s.1688.com/company/pc/factory_search.htm"
            f"?keywords={encoded_keyword}&province={encoded_province}"
        )

    def write_workbook(self, file_path: Path, lead_list: List[Dict], is_phone_masked: bool) -> None:
        Workbook, Font, PatternFill, get_column_letter = self.load_openpyxl()
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "招商线索"
        sheet.append([header_text for _, header_text in EXPORT_HEADER])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="1F2937")
            cell.fill = PatternFill(fill_type="solid", fgColor="E5E7EB")
        for lead in lead_list:
            row = []
            for field_name, _ in EXPORT_HEADER:
                value = lead.get(field_name, "")
                if field_name == "registered_address":
                    value = clean_registered_address_text(value)
                if field_name in ("shop_url", "profile_url"):
                    value = self.format_export_url(lead, field_name)
                if field_name == "phone_source":
                    value = self.format_export_phone_source(lead)
                if field_name == "phone" and is_phone_masked:
                    value = mask_phone(value)
                if field_name == "assignment_status":
                    value = "已分配" if value == ASSIGNMENT_STATUS_ASSIGNED else "未分配"
                row.append(value)
            sheet.append(row)
        for column_index, _ in enumerate(EXPORT_HEADER, start=1):
            sheet.column_dimensions[get_column_letter(column_index)].width = 18
        workbook.save(file_path)

    def load_openpyxl(self):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
            from openpyxl.utils import get_column_letter
        except AttributeError as ex:
            raise RuntimeError(
                "当前 Python 环境的 openpyxl 与 numpy 版本不兼容，请使用 backend\\.venv 启动后端，"
                "或执行 .\\.venv\\Scripts\\python -m pip install -r requirements.txt"
            ) from ex
        return Workbook, Font, PatternFill, get_column_letter
