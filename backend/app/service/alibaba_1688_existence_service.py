import base64
import binascii
import html
import io
import re
import time
import uuid
from dataclasses import dataclass
from datetime import datetime
from html.parser import HTMLParser
from pathlib import Path
from threading import Lock
from typing import Dict, List, Optional
from urllib import error as urllib_error
from urllib import parse as urllib_parse
from urllib import request as urllib_request
from xml.etree import ElementTree

from openpyxl import Workbook, load_workbook

from app.core.config import get_config
from app.service.collector_service import (
    Alibaba1688PublicPageCollector,
    clean_company_name_text,
    sync_playwright,
)


ALIBABA_1688_STATUS_EXISTS = "存在"
ALIBABA_1688_STATUS_NOT_EXISTS = "不存在"
ALIBABA_1688_STATUS_UNCONFIRMED = "未确认"
ALIBABA_1688_JOB_STATUS_PENDING = "pending"
ALIBABA_1688_JOB_STATUS_RUNNING = "running"
ALIBABA_1688_JOB_STATUS_COMPLETED = "completed"
ALIBABA_1688_JOB_STATUS_FAILED = "failed"

ALIBABA_1688_RESULT_HEADER_LIST = [
    "1688反证状态",
    "1688匹配公司名",
    "1688证据链接",
    "1688检索词",
    "1688反证备注",
    "1688反证时间",
]
ALIBABA_1688_EXISTENCE_RESULT_HEADER_LIST = ALIBABA_1688_RESULT_HEADER_LIST
ALIBABA_1688_RESULT_SHEET_NAME_LIST = ["全部结果", "存在", "不存在", "未确认"]
COMPANY_NAME_HEADER_CANDIDATE_LIST = [
    "企业ID",
    "公司名称",
    "企业名称",
    "公司名",
    "单位名称",
]

GUANGDONG_CITY_LIST = [
    "广州",
    "深圳",
    "东莞",
    "佛山",
    "中山",
    "珠海",
    "惠州",
    "汕头",
    "潮州",
    "揭阳",
    "江门",
    "肇庆",
    "湛江",
    "茂名",
    "阳江",
    "清远",
    "韶关",
    "梅州",
    "河源",
    "汕尾",
    "云浮",
]

COMPANY_REGION_PREFIX_LIST = [
    "北京市?",
    "上海市?",
    "天津市?",
    "重庆市?",
    "河北省?",
    "山西省?",
    "辽宁省?",
    "吉林省?",
    "黑龙江省?",
    "江苏省?",
    "浙江省?",
    "安徽省?",
    "福建省?",
    "江西省?",
    "山东省?",
    "河南省?",
    "湖北省?",
    "湖南省?",
    "广东省?",
    "海南省?",
    "四川省?",
    "贵州省?",
    "云南省?",
    "陕西省?",
    "甘肃省?",
    "青海省?",
    "台湾省?",
    "内蒙古(?:自治区)?",
    "广西(?:壮族自治区)?",
    "西藏(?:自治区)?",
    "宁夏(?:回族自治区)?",
    "新疆(?:维吾尔自治区)?",
    "香港(?:特别行政区)?",
    "澳门(?:特别行政区)?",
]

COMPANY_SUFFIX_PATTERN = (
    "股份有限公司|有限责任公司|集团有限公司|有限公司|"
    "分公司|合伙企业|个人独资企业|商行|经营部|工作室|工厂|厂"
)


@dataclass
class Alibaba1688ExistenceResult:
    source_company_name: str
    status: str
    matched_company_name: str
    search_url: str
    search_keyword: str
    remark: str
    verified_at: str


@dataclass
class SearchIndexResult:
    title: str
    url: str
    snippet: str


@dataclass
class Alibaba1688ExistenceJob:
    job_id: str
    file_name: str
    header_list: List
    company_row_list: List[Dict]
    result_list: List[Alibaba1688ExistenceResult]
    status: str = ALIBABA_1688_JOB_STATUS_PENDING
    current_row_index: int = 0
    current_company_name: str = ""
    current_search_keyword: str = ""
    error_message: str = ""
    file_path: str = ""
    download_url: str = ""
    exists_file_path: str = ""
    exists_file_name: str = ""
    exists_download_url: str = ""
    not_exists_file_path: str = ""
    not_exists_file_name: str = ""
    not_exists_download_url: str = ""


class BingSearchResultHtmlParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.result_list: List[SearchIndexResult] = []
        self.current_result: Optional[Dict[str, List[str]]] = None
        self.is_reading_title = False
        self.is_reading_snippet = False

    def handle_starttag(self, tag, attrs):
        attr_map = {name: value or "" for name, value in attrs}
        class_name = attr_map.get("class", "")
        if tag == "li" and "b_algo" in class_name:
            self.flush_current_result()
            self.current_result = {"title": [], "snippet": [], "url": [""]}
            return
        if self.current_result is None:
            return
        if tag == "h2":
            self.is_reading_title = True
        elif tag == "p":
            self.is_reading_snippet = True
        elif tag == "a" and not self.current_result.get("url", [""])[0]:
            self.current_result["url"] = [attr_map.get("href", "")]

    def handle_endtag(self, tag):
        if tag == "h2":
            self.is_reading_title = False
        elif tag == "p":
            self.is_reading_snippet = False

    def handle_data(self, data):
        if self.current_result is None:
            return
        clean_data = re.sub(r"\s+", " ", data or "").strip()
        if not clean_data:
            return
        if self.is_reading_title:
            self.current_result["title"].append(clean_data)
        elif self.is_reading_snippet:
            self.current_result["snippet"].append(clean_data)

    def close(self):
        self.flush_current_result()
        super().close()

    def flush_current_result(self):
        if self.current_result is None:
            return
        title = " ".join(self.current_result.get("title", [])).strip()
        url = (self.current_result.get("url") or [""])[0]
        snippet = " ".join(self.current_result.get("snippet", [])).strip()
        if title or url or snippet:
            self.result_list.append(
                SearchIndexResult(
                    title=html.unescape(title),
                    url=decode_bing_result_url(url),
                    snippet=html.unescape(snippet),
                )
            )
        self.current_result = None


class BingSearchIndexProvider:
    source_name = "Bing/微软搜索"

    def build_search_url(self, query: str) -> str:
        return "https://www.bing.com/search?" + urllib_parse.urlencode({"q": query, "mkt": "zh-CN"})

    def build_rss_search_url(self, query: str) -> str:
        return "https://www.bing.com/search?" + urllib_parse.urlencode(
            {"q": query, "format": "rss", "mkt": "zh-CN"}
        )

    def search(self, query: str, limit: int = 10) -> List[SearchIndexResult]:
        result_limit = max(limit, 1)
        rss_text = self.fetch_url_text(self.build_rss_search_url(query))
        if has_search_access_limit(rss_text):
            raise RuntimeError(f"{self.source_name}出现验证或访问限制")
        rss_result_list = self.parse_rss_result_list(rss_text)
        if rss_result_list:
            return rss_result_list[:result_limit]

        html_text = self.fetch_url_text(self.build_search_url(query))
        if has_search_access_limit(html_text):
            raise RuntimeError(f"{self.source_name}出现验证或访问限制")

        parser = BingSearchResultHtmlParser()
        parser.feed(html_text)
        parser.close()
        return parser.result_list[:result_limit]

    def fetch_url_text(self, search_url: str) -> str:
        request = urllib_request.Request(
            search_url,
            headers={
                "User-Agent": (
                    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                    "AppleWebKit/537.36 (KHTML, like Gecko) "
                    "Chrome/124.0 Safari/537.36"
                ),
                "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8",
            },
        )
        try:
            with urllib_request.urlopen(request, timeout=15) as response:
                return response.read().decode("utf-8", errors="ignore")
        except (urllib_error.URLError, TimeoutError) as ex:
            raise RuntimeError(f"{self.source_name}访问失败：{ex}") from ex

    def parse_rss_result_list(self, rss_text: str) -> List[SearchIndexResult]:
        try:
            root = ElementTree.fromstring(rss_text or "")
        except ElementTree.ParseError:
            return []
        result_list = []
        for item in root.findall("./channel/item"):
            title = item.findtext("title") or ""
            link = item.findtext("link") or ""
            description = item.findtext("description") or ""
            result_list.append(
                SearchIndexResult(
                    title=html.unescape(title).strip(),
                    url=decode_bing_result_url((link or "").strip()),
                    snippet=html.unescape(description).strip(),
                )
            )
        return result_list


class BrowserSearchIndexProvider:
    source_name = "浏览器搜索"

    def __init__(self, config=None, fallback_provider: Optional[BingSearchIndexProvider] = None):
        self.config = config or get_config()
        self.fallback_provider = fallback_provider or BingSearchIndexProvider()
        self.playwright = None
        self.browser = None
        self.page = None

    def build_search_url(self, query: str) -> str:
        return "https://www.bing.com/search?" + urllib_parse.urlencode({"q": query, "mkt": "zh-CN"})

    def search(self, query: str, limit: int = 10) -> List[SearchIndexResult]:
        try:
            page = self.ensure_page()
            page.goto(
                self.build_search_url(query),
                wait_until="domcontentloaded",
                timeout=max(8000, int(getattr(self.config, "web_search_timeout_seconds", 15) * 1000)),
            )
            page.wait_for_timeout(1800)
            page_text = page.locator("body").inner_text(timeout=5000)
            if has_search_access_limit(page_text):
                raise RuntimeError("搜索页出现验证或访问限制")
            anchor_list = page.evaluate(
                """
                () => Array.from(document.querySelectorAll('a[href]')).slice(0, 320).map((anchor) => {
                  const block = anchor.closest('li.b_algo, .b_algo, .result, .vrwrap, article, .algo') || anchor.parentElement;
                  const normalizeText = (value) => (value || '').replace(/\\s+/g, ' ').trim();
                  let href = anchor.getAttribute('href') || '';
                  try {
                    href = new URL(href, location.href).href;
                  } catch (error) {
                    href = anchor.href || '';
                  }
                  return {
                    title: normalizeText(anchor.innerText || anchor.textContent || anchor.getAttribute('aria-label') || ''),
                    url: href,
                    snippet: normalizeText(block ? block.innerText : '')
                  };
                })
                """
            )
            result_list = self.build_result_list_from_anchor_list(anchor_list, limit)
            if result_list:
                return result_list
            return self.fallback_provider.search(query, limit)
        except Exception:
            return self.fallback_provider.search(query, limit)

    def ensure_page(self):
        if self.page is not None:
            return self.page
        if sync_playwright is None:
            raise RuntimeError("Playwright 未安装，无法打开浏览器搜索")
        self.playwright = sync_playwright().start()
        launch_args = {"headless": bool(getattr(self.config, "alibaba_1688_headless", False))}
        browser_channel = getattr(self.config, "alibaba_1688_browser_channel", "")
        if browser_channel:
            launch_args["channel"] = browser_channel
        self.browser = self.playwright.chromium.launch(**launch_args)
        self.page = self.browser.new_page(locale="zh-CN", viewport={"width": 1366, "height": 900})
        return self.page

    def close(self) -> None:
        try:
            if self.browser is not None:
                self.browser.close()
        finally:
            self.browser = None
            self.page = None
            if self.playwright is not None:
                self.playwright.stop()
            self.playwright = None

    @staticmethod
    def build_result_list_from_anchor_list(anchor_list: List[Dict], limit: int = 10) -> List[SearchIndexResult]:
        result_list = []
        seen_url_set = set()
        for anchor in anchor_list or []:
            title = re.sub(r"\s+", " ", str(anchor.get("title") or "")).strip()
            url = decode_bing_result_url(str(anchor.get("url") or "").strip())
            snippet = re.sub(r"\s+", " ", str(anchor.get("snippet") or "")).strip()
            if not url.startswith(("http://", "https://")):
                continue
            dedupe_key = url.split("#", 1)[0].rstrip("/")
            if dedupe_key in seen_url_set:
                continue
            seen_url_set.add(dedupe_key)
            result_list.append(SearchIndexResult(title=title, url=url, snippet=snippet))
            if len(result_list) >= limit:
                break
        return result_list


def normalize_1688_company_name_for_match(company_name: str) -> str:
    clean_name = clean_1688_company_name_preserving_full_name(company_name or "")
    clean_name = re.sub(r"[\u200b\u200c\u200d\ufeff]", "", clean_name)
    clean_name = re.sub(r"\s+", "", clean_name)
    clean_name = clean_name.strip(" \t\r\n,，.。;；:：()（）[]【】")
    if clean_name.startswith("广东省"):
        clean_name = clean_name[3:]
    elif clean_name.startswith("广东"):
        clean_name = clean_name[2:]
    for city_name in GUANGDONG_CITY_LIST:
        city_with_suffix = f"{city_name}市"
        if clean_name.startswith(city_with_suffix):
            return city_name + clean_name[len(city_with_suffix) :]
        if clean_name.startswith(city_name):
            return clean_name
    return clean_name


def clean_1688_company_name_preserving_full_name(company_name: str) -> str:
    raw_clean_name = re.sub(r"\s+", "", company_name or "")
    raw_clean_name = raw_clean_name.strip(" \t\r\n,，.。;；:：()（）[]【】")
    clean_name = clean_company_name_text(company_name or "")
    if not clean_name:
        return raw_clean_name
    if raw_clean_name.startswith(clean_name) and len(raw_clean_name) > len(clean_name):
        return raw_clean_name
    if raw_clean_name and raw_clean_name != clean_name and len(raw_clean_name) <= 100:
        if re.search(r"(有限责任公司|股份有限公司|集团有限公司|有限公司|合伙企业|个人独资企业|分公司|商行|经营部|工作室|工厂|厂)$", raw_clean_name):
            return raw_clean_name
    return clean_name


def decode_bing_result_url(raw_url: str) -> str:
    url = html.unescape(raw_url or "")
    if not url:
        return ""
    parsed_url = urllib_parse.urlparse(url)
    query_param_map = urllib_parse.parse_qs(parsed_url.query)
    bing_encoded_url_list = query_param_map.get("u") or query_param_map.get("url")
    if not bing_encoded_url_list:
        return url
    encoded_url = bing_encoded_url_list[0]
    if encoded_url.startswith("a1"):
        encoded_url = encoded_url[2:]
    try:
        padding = "=" * (-len(encoded_url) % 4)
        decoded_url = base64.urlsafe_b64decode((encoded_url + padding).encode("ascii")).decode("utf-8")
        return decoded_url or url
    except (ValueError, UnicodeDecodeError, binascii.Error):
        return urllib_parse.unquote(encoded_url) or url


def has_search_access_limit(html_text: str) -> bool:
    lowered_text = (html_text or "").lower()
    access_limit_marker_list = [
        "captcha",
        "verify you are human",
        "unusual traffic",
        "访问异常",
        "人机验证",
        "滑块",
        "请完成验证",
    ]
    return any(marker in lowered_text for marker in access_limit_marker_list)


def is_1688_url(url: str) -> bool:
    decoded_url = decode_bing_result_url(url)
    parsed_url = urllib_parse.urlparse(decoded_url)
    host = (parsed_url.netloc or "").lower()
    if ":" in host:
        host = host.split(":", 1)[0]
    return host == "1688.com" or host.endswith(".1688.com")


def is_1688_evidence_url(url: str) -> bool:
    if not is_1688_url(url):
        return False
    decoded_url = decode_bing_result_url(url)
    parsed_url = urllib_parse.urlparse(decoded_url)
    host = (parsed_url.netloc or "").lower().split(":", 1)[0]
    path = (parsed_url.path or "/").lower()
    if host in {"1688.com", "www.1688.com"} and path in {"", "/", "/index.html"}:
        return False
    if any(token in path for token in ("login", "passport", "member", "help", "notice")):
        return False
    return True


def normalize_search_index_text_for_match(text: str) -> str:
    clean_text = re.sub(r"[\u200b\u200c\u200d\ufeff]", "", text or "")
    clean_text = re.sub(r"\s+", "", clean_text)
    clean_text = re.sub(r"[,\.\-_/\\|｜·，。;；:：()（）\[\]【】《》<>「」“”\"'、]", "", clean_text)
    clean_text = clean_text.replace("广东省", "").replace("广东", "")
    for city_name in GUANGDONG_CITY_LIST:
        clean_text = clean_text.replace(f"{city_name}市", city_name)
    return clean_text


def search_index_text_contains_company_name(evidence_text: str, source_company_name: str) -> bool:
    clean_source_name = clean_1688_company_name_preserving_full_name(source_company_name or "")
    if not clean_source_name:
        return False
    normalized_source_name = normalize_search_index_text_for_match(clean_source_name)
    normalized_evidence_text = normalize_search_index_text_for_match(evidence_text or "")
    return bool(normalized_source_name and normalized_source_name in normalized_evidence_text)


def find_enterprise_id_column_index(header_list: List[str]) -> int:
    for index, header in enumerate(header_list or [], start=1):
        if str(header or "").strip() == "企业ID":
            return index
    return 0


def find_company_name_column_index(header_list: List[str], manual_column_index: int = 0) -> int:
    if manual_column_index > 0:
        return manual_column_index if manual_column_index <= len(header_list or []) else 0
    normalized_header_list = [str(header or "").strip() for header in header_list or []]
    for candidate_header in COMPANY_NAME_HEADER_CANDIDATE_LIST:
        for index, header in enumerate(normalized_header_list, start=1):
            if header == candidate_header:
                return index
    return 0


class Alibaba1688ExistenceService:
    def __init__(
        self,
        collector: Optional[Alibaba1688PublicPageCollector] = None,
        config=None,
        search_index_provider: Optional[BingSearchIndexProvider] = None,
    ):
        self.collector = collector or Alibaba1688PublicPageCollector()
        self.config = config or get_config()
        self.search_index_provider = search_index_provider or BrowserSearchIndexProvider(
            self.config,
            fallback_provider=BingSearchIndexProvider(),
        )
        self.job_lock = Lock()
        self.job_map: Dict[str, Alibaba1688ExistenceJob] = {}

    def judge_company_result(
        self,
        source_company_name: str,
        candidate_company_name_list: List[str],
        search_url: str = "",
        search_keyword: str = "",
    ) -> Alibaba1688ExistenceResult:
        verified_at = self.build_verified_at()
        clean_source_name = clean_1688_company_name_preserving_full_name(source_company_name or "")
        if not clean_source_name:
            return Alibaba1688ExistenceResult(
                source_company_name=source_company_name or "",
                status=ALIBABA_1688_STATUS_UNCONFIRMED,
                matched_company_name="",
                search_url=search_url or "",
                search_keyword=search_keyword or "",
                remark="源企业名为空",
                verified_at=verified_at,
            )

        normalized_source_name = normalize_1688_company_name_for_match(clean_source_name)
        for candidate_company_name in candidate_company_name_list:
            clean_candidate_name = clean_1688_company_name_preserving_full_name(candidate_company_name or "")
            if not clean_candidate_name:
                continue
            if normalize_1688_company_name_for_match(clean_candidate_name) == normalized_source_name:
                return Alibaba1688ExistenceResult(
                    source_company_name=clean_source_name,
                    status=ALIBABA_1688_STATUS_EXISTS,
                    matched_company_name=clean_candidate_name,
                    search_url=search_url or "",
                    search_keyword=search_keyword or "",
                    remark="企业名严格匹配",
                    verified_at=verified_at,
                )

        return Alibaba1688ExistenceResult(
            source_company_name=clean_source_name,
            status=ALIBABA_1688_STATUS_NOT_EXISTS,
            matched_company_name="",
            search_url=search_url or "",
            search_keyword=search_keyword or "",
            remark="搜索页候选企业名未严格匹配",
            verified_at=verified_at,
        )

    def read_company_rows_from_excel_bytes(self, excel_bytes: bytes, company_name_column_index: int = 0) -> Dict:
        workbook = load_workbook(io.BytesIO(excel_bytes), data_only=True)
        sheet = workbook.active
        header_list = [cell.value if cell.value is not None else "" for cell in sheet[1]]
        detected_column_index = find_company_name_column_index(header_list, company_name_column_index)
        if detected_column_index <= 0:
            raise ValueError("请选择公司名称列")

        company_row_list = []
        for row_index in range(2, sheet.max_row + 1):
            row_values = [
                sheet.cell(row=row_index, column=column_index).value
                for column_index in range(1, len(header_list) + 1)
            ]
            company_name = row_values[detected_column_index - 1]
            company_row_list.append(
                {
                    "row_index": row_index,
                    "company_name": "" if company_name is None else str(company_name).strip(),
                    "row_values": row_values,
                }
            )
        return {
            "header_list": header_list,
            "enterprise_id_column_index": find_enterprise_id_column_index(header_list),
            "company_name_column_index": detected_column_index,
            "company_row_list": company_row_list,
        }

    def preview_excel_columns_from_bytes(self, excel_bytes: bytes) -> Dict:
        workbook = load_workbook(io.BytesIO(excel_bytes), read_only=True, data_only=True)
        sheet = workbook.active
        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), [])
        header_list = ["" if value is None else str(value).strip() for value in header_row]
        suggested_column_index = find_company_name_column_index(header_list)
        suggested_column_name = (
            header_list[suggested_column_index - 1]
            if suggested_column_index > 0 and suggested_column_index <= len(header_list)
            else ""
        )
        return {
            "headers": header_list,
            "suggested_column_index": suggested_column_index,
            "suggested_column_name": suggested_column_name,
        }

    def write_result_workbook(
        self,
        header_list: List,
        company_row_list: List[Dict],
        result_list: List[Alibaba1688ExistenceResult],
    ) -> bytes:
        workbook = Workbook()
        workbook.active.title = "全部结果"
        for sheet_name in ALIBABA_1688_RESULT_SHEET_NAME_LIST[1:]:
            workbook.create_sheet(sheet_name)
        row_result_pair_list = list(zip(company_row_list, result_list))
        self.write_result_sheet(workbook["全部结果"], header_list, row_result_pair_list)
        for status in (ALIBABA_1688_STATUS_EXISTS, ALIBABA_1688_STATUS_NOT_EXISTS, ALIBABA_1688_STATUS_UNCONFIRMED):
            self.write_result_sheet(
                workbook[status],
                header_list,
                [(row, result) for row, result in row_result_pair_list if result.status == status],
            )
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()

    def write_filtered_result_workbook(
        self,
        header_list: List,
        row_result_pair_list: List,
        sheet_name: str,
    ) -> bytes:
        workbook = Workbook()
        workbook.active.title = sheet_name
        self.write_result_sheet(workbook.active, header_list, row_result_pair_list)
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()

    def write_result_sheet(self, sheet, header_list: List, row_result_pair_list: List) -> None:
        sheet.append(list(header_list) + ALIBABA_1688_RESULT_HEADER_LIST)
        for company_row, result in row_result_pair_list:
            sheet.append(
                list(company_row.get("row_values", []))
                + [
                    result.status,
                    result.matched_company_name,
                    result.search_url,
                    result.search_keyword,
                    result.remark,
                    result.verified_at,
                ]
            )

    def extract_candidate_company_name_list(self, raw_item_list: List[Dict]) -> List[str]:
        candidate_name_list = []
        seen_name_set = set()
        guangdong_city_pattern = rf"(?:广东省|广东)?(?:{'|'.join(GUANGDONG_CITY_LIST)})(?:市)?"
        region_prefix_pattern = rf"(?:{'|'.join(COMPANY_REGION_PREFIX_LIST)})"
        candidate_pattern_list = [
            re.compile(
                rf"({guangdong_city_pattern}[\u4e00-\u9fffA-Za-z0-9（）()·\-]{{1,80}}"
                rf"(?:{COMPANY_SUFFIX_PATTERN}))"
            ),
            re.compile(
                rf"({region_prefix_pattern}[\u4e00-\u9fffA-Za-z0-9（）()·\-]{{1,80}}"
                rf"(?:{COMPANY_SUFFIX_PATTERN}))"
            ),
        ]
        for raw_item in raw_item_list or []:
            text = ""
            if isinstance(raw_item, dict):
                text = " ".join(
                    str(raw_item.get(field_name) or "")
                    for field_name in ("company_name", "title", "text", "parentText")
                )
            else:
                text = str(raw_item or "")
            clean_text = re.sub(r"\s+", " ", text).strip()
            for pattern in candidate_pattern_list:
                for match in pattern.findall(clean_text):
                    candidate_name = clean_1688_company_name_preserving_full_name(match)
                    if candidate_name and candidate_name not in seen_name_set:
                        seen_name_set.add(candidate_name)
                        candidate_name_list.append(candidate_name)
        return candidate_name_list

    def verify_company_on_1688_search_page(self, company_name: str, page=None) -> Alibaba1688ExistenceResult:
        keyword_list = self.build_1688_existence_search_keyword_list(company_name)
        if not keyword_list:
            return self.judge_company_result(company_name, [], "", "")

        last_search_url = ""
        last_search_keyword = ""
        candidate_company_name_list = []
        has_unparsed_result_item = False
        for keyword in keyword_list:
            last_search_keyword = keyword
            search_url_list = self.collector.build_search_url_list(keyword, 1)
            for search_url in search_url_list:
                last_search_url = search_url
                self.collector.goto_public_page(page, search_url)
                self.collector.scroll_result_page_to_load(page)
                raw_item_list = self.collector.extract_raw_item_list(page)
                if raw_item_list:
                    has_unparsed_result_item = True
                candidate_company_name_list.extend(self.extract_candidate_company_name_list(raw_item_list))
                if candidate_company_name_list:
                    result = self.judge_company_result(
                        company_name,
                        candidate_company_name_list,
                        last_search_url,
                        keyword,
                    )
                    if result.status == ALIBABA_1688_STATUS_EXISTS:
                        return result
            if candidate_company_name_list:
                break
        if not candidate_company_name_list and has_unparsed_result_item:
            return Alibaba1688ExistenceResult(
                source_company_name=clean_1688_company_name_preserving_full_name(company_name or ""),
                status=ALIBABA_1688_STATUS_UNCONFIRMED,
                matched_company_name="",
                search_url=last_search_url,
                search_keyword=last_search_keyword,
                remark="搜索页有结果但未解析到企业名",
                verified_at=self.build_verified_at(),
            )
        return self.judge_company_result(
            company_name,
            candidate_company_name_list,
            last_search_url,
            last_search_keyword,
        )

    def verify_company_on_public_search_index(self, company_name: str) -> Alibaba1688ExistenceResult:
        query_list = self.build_public_search_index_query_list(company_name)
        if not query_list:
            return self.judge_company_result(company_name, [], "", "")

        last_query = query_list[0]
        last_search_page_url = self.search_index_provider.build_search_url(last_query)
        for query in query_list:
            last_query = query
            last_search_page_url = self.search_index_provider.build_search_url(query)
            search_result_list = self.search_index_provider.search(query, limit=10)
            result = self.judge_public_search_index_result(
                company_name,
                search_result_list,
                last_search_page_url,
                query,
            )
            if result.status == ALIBABA_1688_STATUS_EXISTS:
                return result
        return self.judge_public_search_index_result(company_name, [], last_search_page_url, last_query)

    def judge_public_search_index_result(
        self,
        source_company_name: str,
        search_result_list: List[SearchIndexResult],
        search_page_url: str,
        search_keyword: str,
    ) -> Alibaba1688ExistenceResult:
        verified_at = self.build_verified_at()
        clean_source_name = clean_1688_company_name_preserving_full_name(source_company_name or "")
        if not clean_source_name:
            return Alibaba1688ExistenceResult(
                source_company_name=source_company_name or "",
                status=ALIBABA_1688_STATUS_UNCONFIRMED,
                matched_company_name="",
                search_url=search_page_url or "",
                search_keyword=search_keyword or "",
                remark="源企业名为空",
                verified_at=verified_at,
            )

        has_1688_result = False
        for search_result in search_result_list or []:
            if not is_1688_evidence_url(search_result.url):
                continue
            has_1688_result = True
            title = re.sub(r"\s+", " ", search_result.title or "").strip()
            return Alibaba1688ExistenceResult(
                source_company_name=clean_source_name,
                status=ALIBABA_1688_STATUS_EXISTS,
                matched_company_name=clean_source_name,
                search_url=search_result.url or search_page_url or "",
                search_keyword=search_keyword or "",
                remark=f"{self.search_index_provider.source_name}搜索页出现1688链接：{title[:80]}",
                verified_at=verified_at,
            )

        if has_1688_result:
            remark = f"{self.search_index_provider.source_name}发现1688结果，但企业名未严格匹配"
        else:
            remark = f"{self.search_index_provider.source_name}前10条未发现1688公开索引"
        return Alibaba1688ExistenceResult(
            source_company_name=clean_source_name,
            status=ALIBABA_1688_STATUS_NOT_EXISTS,
            matched_company_name="",
            search_url=search_page_url or "",
            search_keyword=search_keyword or "",
            remark=remark,
            verified_at=verified_at,
        )

    def build_1688_existence_search_keyword_list(self, company_name: str) -> List[str]:
        keyword_list = []
        for keyword in (
            clean_1688_company_name_preserving_full_name(company_name or ""),
            normalize_1688_company_name_for_match(company_name or ""),
        ):
            if keyword and keyword not in keyword_list:
                keyword_list.append(keyword)
        return keyword_list

    def build_public_search_index_query_list(self, company_name: str) -> List[str]:
        query_list = []
        for keyword in self.build_1688_existence_search_keyword_list(company_name):
            query = f"{keyword} 阿里巴巴"
            if query not in query_list:
                query_list.append(query)
        return query_list

    def decode_excel_base64(self, excel_base64: str, file_name: str) -> bytes:
        if not str(file_name or "").lower().endswith(".xlsx"):
            raise ValueError("仅支持.xlsx文件")
        try:
            return base64.b64decode(excel_base64, validate=True)
        except (binascii.Error, ValueError) as ex:
            raise ValueError("Excel base64内容无效") from ex

    def preview_excel_columns_base64(self, excel_base64: str, file_name: str) -> Dict:
        excel_bytes = self.decode_excel_base64(excel_base64, file_name)
        return self.preview_excel_columns_from_bytes(excel_bytes)

    def verify_excel_base64(self, excel_base64: str, file_name: str, company_name_column_index: int = 0) -> Dict:
        excel_bytes = self.decode_excel_base64(excel_base64, file_name)
        excel_data = self.read_company_rows_from_excel_bytes(
            excel_bytes,
            company_name_column_index=company_name_column_index,
        )
        company_row_list = excel_data["company_row_list"]
        result_list = []

        for index, company_row in enumerate(company_row_list):
            result_list.append(self.verify_company_row_with_fallback(company_row, page=None))
            if index < len(company_row_list) - 1:
                time.sleep(max(getattr(self.config, "alibaba_1688_request_interval_seconds", 1.0), 1.0))

        return self.write_result_exports(excel_data["header_list"], company_row_list, result_list, file_name)

    def write_result_exports(
        self,
        header_list: List,
        company_row_list: List[Dict],
        result_list: List[Alibaba1688ExistenceResult],
        file_name: str,
    ) -> Dict:
        row_result_pair_list = list(zip(company_row_list, result_list))
        result_workbook_bytes = self.write_result_workbook(header_list, company_row_list, result_list)
        export_file_path = self.build_export_file_path(file_name, "1688_existence")
        export_file_path.write_bytes(result_workbook_bytes)
        exists_export_file_path = self.build_export_file_path(file_name, "1688_exists")
        exists_export_file_path.write_bytes(
            self.write_filtered_result_workbook(
                header_list,
                [
                    (company_row, result)
                    for company_row, result in row_result_pair_list
                    if result.status == ALIBABA_1688_STATUS_EXISTS
                ],
                ALIBABA_1688_STATUS_EXISTS,
            )
        )
        not_exists_export_file_path = self.build_export_file_path(file_name, "1688_not_exists")
        not_exists_export_file_path.write_bytes(
            self.write_filtered_result_workbook(
                header_list,
                [
                    (company_row, result)
                    for company_row, result in row_result_pair_list
                    if result.status == ALIBABA_1688_STATUS_NOT_EXISTS
                ],
                ALIBABA_1688_STATUS_NOT_EXISTS,
            )
        )
        return self.build_summary(
            company_row_list,
            result_list,
            export_file_path,
            exists_export_file_path,
            not_exists_export_file_path,
        )

    def create_verification_job(
        self,
        excel_base64: str,
        file_name: str,
        company_name_column_index: int = 0,
    ) -> Dict:
        excel_bytes = self.decode_excel_base64(excel_base64, file_name)
        excel_data = self.read_company_rows_from_excel_bytes(
            excel_bytes,
            company_name_column_index=company_name_column_index,
        )
        job = Alibaba1688ExistenceJob(
            job_id=uuid.uuid4().hex,
            file_name=file_name,
            header_list=excel_data["header_list"],
            company_row_list=excel_data["company_row_list"],
            result_list=[],
        )
        with self.job_lock:
            self.job_map[job.job_id] = job
            return self.build_job_snapshot(job)

    def run_verification_job(self, job_id: str) -> None:
        with self.job_lock:
            job = self.get_verification_job(job_id)
            job.status = ALIBABA_1688_JOB_STATUS_RUNNING
            job.error_message = ""
        try:
            for index, company_row in enumerate(job.company_row_list):
                company_name = company_row.get("company_name", "")
                keyword_list = self.build_public_search_index_query_list(company_name)
                with self.job_lock:
                    job.current_row_index = company_row.get("row_index", index + 2)
                    job.current_company_name = company_name
                    job.current_search_keyword = keyword_list[0] if keyword_list else company_name
                result = self.verify_company_row_with_fallback(company_row, page=None)
                with self.job_lock:
                    job.result_list.append(result)
                    job.current_search_keyword = result.search_keyword
                if index < len(job.company_row_list) - 1:
                    time.sleep(max(getattr(self.config, "alibaba_1688_request_interval_seconds", 1.0), 1.0))
            summary = self.write_result_exports(job.header_list, job.company_row_list, job.result_list, job.file_name)
            with self.job_lock:
                job.status = ALIBABA_1688_JOB_STATUS_COMPLETED
                job.current_row_index = 0
                job.current_company_name = ""
                job.current_search_keyword = ""
                job.file_path = summary.get("file_path", "")
                job.download_url = summary.get("download_url", "")
                job.exists_file_path = summary.get("exists_file_path", "")
                job.exists_file_name = summary.get("exists_file_name", "")
                job.exists_download_url = summary.get("exists_download_url", "")
                job.not_exists_file_path = summary.get("not_exists_file_path", "")
                job.not_exists_file_name = summary.get("not_exists_file_name", "")
                job.not_exists_download_url = summary.get("not_exists_download_url", "")
        except Exception as ex:
            with self.job_lock:
                job.status = ALIBABA_1688_JOB_STATUS_FAILED
                job.error_message = str(ex)[:300]
        finally:
            close_provider = getattr(self.search_index_provider, "close", None)
            if callable(close_provider):
                try:
                    close_provider()
                except Exception:
                    pass

    def get_verification_job_snapshot(self, job_id: str) -> Dict:
        with self.job_lock:
            return self.build_job_snapshot(self.get_verification_job(job_id))

    def get_verification_job(self, job_id: str) -> Alibaba1688ExistenceJob:
        job = self.job_map.get(job_id)
        if job is None:
            raise KeyError(job_id)
        return job

    def build_job_snapshot(self, job: Alibaba1688ExistenceJob) -> Dict:
        result_list = list(job.result_list)
        processed_count = len(result_list)
        return {
            "job_id": job.job_id,
            "status": job.status,
            "processed_count": processed_count,
            "current_row_index": job.current_row_index,
            "current_company_name": job.current_company_name,
            "current_search_keyword": job.current_search_keyword,
            "error_message": job.error_message,
            "file_name": Path(job.file_name).name,
            "file_path": job.file_path,
            "download_url": job.download_url,
            "exists_file_name": job.exists_file_name,
            "exists_file_path": job.exists_file_path,
            "exists_download_url": job.exists_download_url,
            "not_exists_file_name": job.not_exists_file_name,
            "not_exists_file_path": job.not_exists_file_path,
            "not_exists_download_url": job.not_exists_download_url,
            "total_count": len(job.company_row_list),
            "exists_count": sum(1 for result in result_list if result.status == ALIBABA_1688_STATUS_EXISTS),
            "not_exists_count": sum(1 for result in result_list if result.status == ALIBABA_1688_STATUS_NOT_EXISTS),
            "unconfirmed_count": sum(1 for result in result_list if result.status == ALIBABA_1688_STATUS_UNCONFIRMED),
            "result_items": self.build_result_item_list(job.company_row_list[:processed_count], result_list),
        }

    def verify_company_row_with_fallback(self, company_row: Dict, page=None) -> Alibaba1688ExistenceResult:
        company_name = company_row.get("company_name", "")
        try:
            return self.verify_company_on_public_search_index(company_name)
        except Exception as ex:
            keyword_list = self.build_public_search_index_query_list(company_name)
            return Alibaba1688ExistenceResult(
                source_company_name=company_name,
                status=ALIBABA_1688_STATUS_UNCONFIRMED,
                matched_company_name="",
                search_url="",
                search_keyword=keyword_list[0] if keyword_list else clean_1688_company_name_preserving_full_name(company_name),
                remark=str(ex)[:180],
                verified_at=self.build_verified_at(),
            )

    def should_use_injected_collector_without_browser(self) -> bool:
        return not isinstance(self.collector, Alibaba1688PublicPageCollector)

    def build_export_file_path(self, file_name: str, export_suffix: str = "1688_existence") -> Path:
        export_dir = Path(self.config.export_dir)
        export_dir.mkdir(parents=True, exist_ok=True)
        stem = Path(file_name).stem or "1688_existence"
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        base_file_path = export_dir / f"{stem}_{export_suffix}_{timestamp}.xlsx"
        if not base_file_path.exists():
            return base_file_path
        for index in range(1, 1000):
            candidate_file_path = export_dir / f"{stem}_{export_suffix}_{timestamp}_{index}.xlsx"
            if not candidate_file_path.exists():
                return candidate_file_path
        return export_dir / f"{stem}_{export_suffix}_{datetime.now().strftime('%Y%m%d_%H%M%S_%f')}.xlsx"

    def build_summary(
        self,
        company_row_list: List[Dict],
        result_list: List[Alibaba1688ExistenceResult],
        export_file_path: Path,
        exists_export_file_path: Optional[Path] = None,
        not_exists_export_file_path: Optional[Path] = None,
    ) -> Dict:
        exists_export_file_path = exists_export_file_path or export_file_path
        not_exists_export_file_path = not_exists_export_file_path or export_file_path
        return {
            "total_count": len(company_row_list),
            "exists_count": sum(1 for result in result_list if result.status == ALIBABA_1688_STATUS_EXISTS),
            "not_exists_count": sum(1 for result in result_list if result.status == ALIBABA_1688_STATUS_NOT_EXISTS),
            "unconfirmed_count": sum(1 for result in result_list if result.status == ALIBABA_1688_STATUS_UNCONFIRMED),
            "result_items": self.build_result_item_list(company_row_list, result_list),
            "file_path": str(export_file_path),
            "file_name": export_file_path.name,
            "download_url": f"/api/exports/{export_file_path.name}",
            "exists_file_path": str(exists_export_file_path),
            "exists_file_name": exists_export_file_path.name,
            "exists_download_url": f"/api/exports/{exists_export_file_path.name}",
            "not_exists_file_path": str(not_exists_export_file_path),
            "not_exists_file_name": not_exists_export_file_path.name,
            "not_exists_download_url": f"/api/exports/{not_exists_export_file_path.name}",
        }

    def build_result_item_list(
        self,
        company_row_list: List[Dict],
        result_list: List[Alibaba1688ExistenceResult],
    ) -> List[Dict]:
        item_list = []
        for company_row, result in zip(company_row_list, result_list):
            item_list.append(
                {
                    "row_index": company_row.get("row_index", 0),
                    "company_name": company_row.get("company_name", ""),
                    "status": result.status,
                    "matched_company_name": result.matched_company_name,
                    "search_url": result.search_url,
                    "search_keyword": result.search_keyword,
                    "remark": result.remark,
                    "verified_at": result.verified_at,
                }
            )
        return item_list

    def build_verified_at(self) -> str:
        return datetime.now().strftime("%Y-%m-%d %H:%M:%S")
